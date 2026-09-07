"""AGUILAR2025 -- Nature Communications Source Data (41467_2025_64427_MOESM6).

Mouse efficacy studies of BPaL-type tuberculosis regimens plus an in vitro
concentration series for JNJ-4052, JNJ-2901 and telacebec.

READ THIS BEFORE USING THE NUMBERS
  - Every "Lung bacterial burden" sheet is headed "Log10 CFU counts/lung" (or
    "Log10 CFU counts"). These are IN VIVO organ burdens: one whole mouse lung
    per value, not a millilitre. cfu_per_ml is filled with 10**value because the
    schema has no per-organ column, and `readout` on every row says so. Do not
    read them as concentrations.
  - Values of exactly 0 (log10) are frequent. They mean 1 CFU/lung as written.
    The deposit states no limit of detection anywhere, so nothing is marked
    censored and floor_cfu_per_ml is blank. Two repeated constants deserve a
    reader's attention, and they do NOT point the same way -- counted from the
    workbook, not assumed:
      * 0.4 (log10) fills every one of the fifteen values in Supplementary
        Table 12's three fully suppressed columns (CZT/BCZ/BCZT 8wks) and is
        the smallest number on that sheet. It behaves like a FLOOR.
      * 4.15 (log10) occurs sixty times in Supplementary Table 4 and is always
        the MAXIMUM of its column, never exceeded (e.g. Study A 'BPaMJ 8wk
        +12wk R' is 4.15 in all fifteen mice; Study B 'BPaJ 8wk +12wk' in ten
        of fifteen, the rest below it). It behaves like a CEILING -- a cap on
        countable relapse, not a detection limit.
    Neither is declared by the deposit, so neither is acted on: floor_cfu_per_ml
    stays blank and the values are returned exactly as deposited.
  - The organism is not named in this workbook. Strain labels H37Rv, N1283 and
    Mtb-N1283 appear on Fig. 3b-e and are kept in `strain`; organism stays blank.
  - The regimen codes (BPaL, BPaMZ, CZ, HRZE, ...) are the source's own labels
    for multi-drug combinations. They are put in `drug` verbatim, not expanded.
  - Durations come from the column headers ("BPaL 8wk", "BPaL/BPa 16wks +
    16wks") and are converted to hours from day 0. Composite labels are summed
    and the verbatim label is kept in `arm` and `notes`. Negative headers
    ("Day -13", "-2wks") get a blank time because the schema forbids one.
  - 'SoT' (Fig. 3e) is given no time by the deposit and is left with a blank
    time_h.

THE 'Supp. Table 4' SHEET HOLDS TWO TABLES, NOT ONE. Rows 1-18 are
"Supplementary Table 4: Lung bacterial burden (Study A)" with its header on row
3; rows 23-40 are a SECOND table, "Supplementary Table 4: Lung bacterial burden
(Study B)", with its own header on row 25 and different regimen columns
(BPaJ, 12wk arms). Every sheet here is therefore parsed block by block: a title
in column A starts a block, its header is two rows below it, and its data run to
the row before the next title. Reading the sheet as one table -- which an
earlier version of this reader did -- put Study B's fifteen rows of mice under
Study A's column headers and gave them Study A's durations.

DUPLICATES NOT RE-READ: Fig. 1b, Fig. 1c, Fig. 1d and Fig. 2 reproduce, to the
digit, subsets of the supplementary tables, which also carry the durations the
figures omit. Verified column by column:
  Fig. 1b  = Supp. Table 4 (Study A): SoTX = 'Day 0' C4:C8, BPaL = 'BPaL 8wk'
             J4:J8, and so on for BPaM/BPaMZ/BPaMJ/BPaC/BPaCJ.
  Fig. 1c  = Supp. Table 4 (Study B), the block on rows 25-40: SoTX = 'Day 0'
             C26:C30, BPaL = 'BPaL 8wk' D26:D30, BPa901 = 'BPaJ 8wk' E26:E30,
             BPaC = 'BPaC 8w' H26:H30, BPaCJ = 'BPaCJ 8wk' I26:I30 -- all five
             columns identical.
  Fig. 1d  = Supp. Table 11 (Study D), rounded to 2 dp.
  Fig. 2   = Supp. Table 12 (Study E), identical.
Only the supplementary tables are read, so no mouse is counted twice.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.ingest import COLUMNS

BOOK = "_unpacked/PMC12546632_supplementary/41467_2025_64427_MOESM6_ESM.xlsx"

NOT_READ = (
    "Fig. 1b / 1c / 1d / Fig. 2 duplicate Supp. Tables 4 (Study A block), 4 "
    "(Study B block), 11 and 12 respectively and are skipped; "
    "Fig. 3a is average MIC90 in nM; Supp. Fig. 2 is a mouse survival table "
    "(day of death), not a count; Supp. Table 8's regimen columns carry no "
    "per-column duration -- the sheet title's '16wks Tx + 12wks' is used and "
    "flagged; 41467_2025_64427_MOESM3 (SD1) is a table of Fisher/Dunnett "
    "p-values with no readings in it"
)

LUNG = ("log10 CFU per mouse lung (the source's heading is 'Log10 CFU counts"
        "/lung'; the denominator is a whole lung, NOT a mL)")
LUNG_FIG = ("log10 CFU per mouse lung (the sheet's heading is 'Lung bacterial "
            "burden ... - Log10 CFU counts'; it says lung but writes no "
            "denominator on the number itself -- NOT a mL)")
INVITRO = ("CFU (the source's heading is 'CFU counts'; it does not state the "
           "denominator)")

# Sheets of per-lung log10 counts. Each is split into blocks by _blocks();
# the study label is read out of each block's OWN title, never assumed.
SHEETS_LOG = ("Supp. Table 4", "Supp. Table 8", "Supp. Table 11",
              "Supp. Table 12")


def _txt(v) -> str:
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def _num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return np.nan


def _blocks(ws):
    """Yield (title, header_row, first_data_row, last_data_row) per table.

    A block starts at a row with text in column A. Its header is two rows
    below; its data run to the row before the next such title. 'Supp. Table 4'
    holds two of these -- Study A on rows 1-18 and Study B on rows 23-40 --
    and reading it as one table silently reassigns Study B's mice to Study A's
    regimen columns.
    """
    titles = [r for r in range(1, ws.max_row + 1)
              if _txt(ws.cell(row=r, column=1).value)]
    for i, r in enumerate(titles):
        end = titles[i + 1] - 1 if i + 1 < len(titles) else ws.max_row
        yield _txt(ws.cell(row=r, column=1).value), r + 2, r + 3, end


def parse_header(h: str) -> tuple[str, float, str]:
    """(regimen, hours from day 0, note) for one column header."""
    s = _txt(h)
    low = s.lower()
    if re.match(r"^day\s*0$", low) or low == "day 0":
        return "", 0.0, "the header is the source's 'Day 0'"
    if re.match(r"^day\s*-\s*\d", low) or re.match(r"^-\s*\d+\s*wk", low):
        return "", np.nan, (f"header '{s}' is BEFORE day 0 (negative time); "
                            f"the schema forbids negative time_h, so it is blank")
    weeks = [float(x) for x in re.findall(r"(\d+(?:\.\d+)?)\s*w(?:ks?)?\b", low)]
    regimen = re.sub(r"\d+(?:\.\d+)?\s*w(?:ks?)?\b", " ", s, flags=re.I)
    regimen = re.sub(r"\+|\bR\b", " ", regimen).strip(" -+")
    if weeks:
        why = f"header '{s}'; weeks converted to hours"
        if len(weeks) > 1:
            why += (f"; the header names {len(weeks)} consecutive periods "
                    f"({' + '.join(str(w) for w in weeks)} weeks) and time_h is "
                    f"their sum from day 0")
        if re.search(r"\d\s*w(?![ks])", low):
            why += ("; this header is written '<n>w' where its siblings "
                    "say '<n>wk' -- read as weeks")
        return regimen, sum(weeks) * 7 * 24.0, why
    return regimen, np.nan, f"header '{s}' gives no time, so time_h is blank"


def read(d: Path) -> pd.DataFrame:
    wb = load_workbook(d / BOOK, data_only=True)
    rows: list[dict] = []

    def add(**kw):
        rows.append({"source_file": BOOK, "organism": "", "strain": "", **kw})

    # ---- Supplementary Tables 4, 8, 11, 12 : log10 CFU per lung ----------
    # Every sheet is parsed BLOCK BY BLOCK. 'Supp. Table 4' carries two.
    for sheet in SHEETS_LOG:
        ws = wb[sheet]
        for title, hrow, first, last in _blocks(ws):
            m = re.search(r"\((Stud(?:y|ies)\s+[A-Z])\)", title, re.I)
            study = m.group(1) if m else ""
            fallback = ""
            m = re.search(r"(\d+)\s*wks?\s*Tx\s*\+\s*(\d+)\s*wks?", title, re.I)
            if m:
                fallback = m.group(0)
            heads = {c: _txt(ws.cell(row=hrow, column=c).value)
                     for c in range(2, ws.max_column + 1)}
            if not any(heads.values()):
                continue
            for r in range(first, last + 1):
                for c, h in heads.items():
                    if not h:
                        continue
                    val = _num(ws.cell(row=r, column=c).value)
                    if np.isnan(val):
                        continue
                    regimen, t, why = parse_header(h)
                    if np.isnan(t) and fallback and regimen and \
                            "before day 0" not in why:
                        wk = [float(x) for x in re.findall(r"(\d+)", fallback)]
                        t = sum(wk) * 7 * 24.0
                        why = (f"header '{h}' carries no duration; the block "
                               f"title says '{fallback}', so time_h is their "
                               f"sum from day 0 -- this comes from the TITLE, "
                               f"not from the column")
                    add(sheet=sheet, drug=regimen,
                        arm=f"{h} [{study}]" if study else h,
                        replicate=f"r{r}", time_h=t,
                        cfu_per_ml=float(10.0 ** val), readout=LUNG,
                        notes=f"{title} | {why} | value as deposited is "
                              f"log10 = {val!r}; regimen codes are the source's "
                              f"own abbreviations and are not expanded; each "
                              f"column is an independent group of mice "
                              f"sacrificed at that time -- rows are NOT the "
                              f"same animal across columns; replicate is the "
                              f"spreadsheet row, which identifies a mouse only "
                              f"within one column")

    # ---- Fig. 3b / 3c / 3d : in vitro concentration series ---------------
    for sheet in ("Fig. 3b", "Fig. 3c", "Fig. 3d"):
        ws = wb[sheet]
        title = _txt(ws["A1"].value)
        m = re.search(r"for\s+(.+?)\s*-\s*CFU counts", title)
        drug = m.group(1).strip() if m else ""
        starts = [c for c in range(1, ws.max_column + 1)
                  if _txt(ws.cell(row=4, column=c).value) == "Conc (nM)"]
        for s0 in starts:
            strain = _txt(ws.cell(row=3, column=s0 + 1).value)
            c = s0 + 1
            concs = {}
            while c <= ws.max_column and _txt(ws.cell(row=4, column=c).value):
                if _txt(ws.cell(row=4, column=c).value) == "Conc (nM)":
                    break
                concs[c] = _txt(ws.cell(row=4, column=c).value)
                c += 1
            for r in range(5, ws.max_row + 1):
                for cc, conc in concs.items():
                    val = _num(ws.cell(row=r, column=cc).value)
                    if np.isnan(val):
                        continue
                    cn = _num(float(conc)) if re.match(r"^[\d.]+$", conc) else np.nan
                    add(sheet=sheet, strain=strain,
                        drug="" if cn == 0 else drug,
                        concentration=cn, conc_unit="nM",
                        arm=f"{strain} {conc} nM [{drug}]",
                        replicate=f"rep{r - 4}", time_h=np.nan,
                        cfu_per_ml=val, readout=INVITRO,
                        notes=f"{title} | the sheet gives no exposure time for "
                              f"this block, so time_h is blank")

    # ---- Fig. 3e : Study F, log10 CFU, strain per row --------------------
    ws = wb["Fig. 3e"]
    title = _txt(ws["A1"].value)
    groups, cur = {}, ""
    for c in range(3, ws.max_column + 1):
        v = _txt(ws.cell(row=3, column=c).value)
        if v:
            cur = v
        groups[c] = cur
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 3:
            v = _txt(ws.cell(row=3, column=rng.min_col).value)
            for c in range(rng.min_col, rng.max_col + 1):
                groups[c] = v
    for r in range(4, ws.max_row + 1):
        strain = _txt(ws.cell(row=r, column=2).value)
        if not strain:
            continue
        n = {}
        for c in range(3, ws.max_column + 1):
            val = _num(ws.cell(row=r, column=c).value)
            arm = groups.get(c, "")
            if np.isnan(val) or not arm:
                continue
            n[arm] = n.get(arm, 0) + 1
            add(sheet="Fig. 3e", strain=strain,
                drug="" if arm.upper().startswith("SOT") or
                     arm.lower() == "vehicle" else arm,
                arm=f"{arm} [Study F]", replicate=f"{strain} row{n[arm]}",
                time_h=np.nan, cfu_per_ml=float(10.0 ** val),
                readout=LUNG_FIG,
                notes=f"{title} | no treatment duration is given on this "
                      f"sheet, so time_h is blank | log10 = {val!r}")

    df = pd.DataFrame(rows)
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = np.nan if c in ("concentration", "time_h", "colonies",
                                    "dilution", "plated_volume_ul",
                                    "cfu_per_ml", "log10_cfu_per_ml",
                                    "floor_cfu_per_ml") else ""
    df["floor_basis"] = (
        "no plated volume, dilution or limit of detection appears anywhere in "
        "this workbook. Two undeclared repeated constants, counted from the "
        "file: 0.4 fills all 15 values of Supp. Table 12's three suppressed "
        "columns and is the smallest number on that sheet (behaves like a "
        "floor); 4.15 occurs 60 times in Supp. Table 4 and is always its "
        "column's MAXIMUM, never exceeded (behaves like a ceiling, not a "
        "floor). Neither is stated, so neither is used")
    return df[COLUMNS]
