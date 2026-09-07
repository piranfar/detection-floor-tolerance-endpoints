"""WALTER2025 -- Nature Communications Source Data (41467_2025_56146_MOESM13).

BTZ-043 (and isoniazid) chemotherapy in mice. The bacterial readings in this
deposit are IN VIVO organ burdens, not broth counts, and that governs how they
are returned.

READ THIS BEFORE USING THE NUMBERS
  - Sheets 1a&b, 1c, 3a and 4a are headed "pulmonary CFU counts / mouse" -- one
    whole mouse lung per value. Sheet 4b is headed "CFU / mm3 (MBLA)". Neither
    denominator is a millilitre. The values are placed in cfu_per_ml because the
    schema has no per-organ column, and `readout` on every row says which
    denominator the source actually used. Do not read these as concentrations.
  - The workbook never names the organism. It is blank on every row.
  - No plated volume, dilution or limit of detection appears anywhere, so
    floor_cfu_per_ml is blank throughout.
  - Time is given as weeks or days relative to therapy start and converted to
    hours. Two labels cannot be turned into a single time and are left blank:
    "3 weeks before therapy start (-3)", which is negative, and "4-6 weeks of
    therapy", which is a range. Both keep their verbatim label in `arm`/`notes`.
  - Sheet 3a labels its baseline column simply "before therapy" with no offset;
    it is taken as t = 0 (start of exposure) and the note says so.

NOT READ: sheets 2a&b, 5f, 6a, 6b, S2a-d, S3a, S3c and S8 of this workbook, and
the whole of MOESM4-11, are LC-MS/MS drug concentrations, Caco-2 permeability or
imaging intensity profiles -- no bacterial counts in any of them.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.ingest import COLUMNS

BOOK = "_unpacked/PMC11742723_supplementary/41467_2025_56146_MOESM13_ESM.xlsx"

NOT_READ = (
    "2a&b, 5f, 6a, 6b, S2a-d, S3a, S3c, S8 in this workbook and all of "
    "MOESM4-MOESM11 are LC-MS/MS drug concentrations, Caco-2 permeability or "
    "imaging intensity, not bacterial counts"
)

LUNG = ("CFU per mouse lung (the source's heading is 'pulmonary CFU counts / "
        "mouse'; the denominator is a whole lung, NOT a mL)")
MBLA = ("CFU per mm3 by MBLA (the source's heading is 'CFU / mm3 (MBLA)'; "
        "NOT per mL)")

_NONE = ("before therapy", "before Therapy", "no treatmen (NT)",
         "no treatment (NT)", "vehicle", "SoT")


def _txt(v) -> str:
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def _num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return np.nan


def _filled(ws) -> dict[tuple[int, int], object]:
    """Cell values with every merged range flooded from its top-left cell."""
    grid = {(c.row, c.column): c.value
            for row in ws.iter_rows() for c in row if c.value is not None}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(row=rng.min_row, column=rng.min_col).value
        if v is None:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid.setdefault((r, c), v)
    return grid


def parse_arm(label: str) -> tuple[str, float, str]:
    """(drug, dose, unit) from a treatment-group label, or blanks."""
    s = _txt(label)
    if not s or s.lower() in {x.lower() for x in _NONE}:
        return "", np.nan, ""
    drug = ""
    for name in ("BTZ-043", "INH"):
        if name in s:
            drug = name
            break
    m = re.search(r"(\d+(?:\.\d+)?)\s*mg\s*/\s*kg", s)
    dose = float(m.group(1)) if m else np.nan
    return drug, dose, "mg/kg" if m else ""


def parse_time(label: str) -> tuple[float, str]:
    """(hours since therapy start, why it is blank if it is)."""
    s = _txt(label)
    low = s.lower()
    if "before therapy" in low and "week" in low:
        return np.nan, (f"source label '{s}' is BEFORE therapy start "
                        f"(negative time); the schema forbids negative time_h, "
                        f"so it is left blank")
    if low.startswith("therapy start"):
        return 0.0, "therapy start (t = 0) as the source labels it"
    if low == "before therapy":
        return 0.0, ("source says only 'before therapy' with no offset; taken "
                     "as t = 0, the start of exposure")
    m = re.match(r"^\s*(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)\s*w", low)
    if m:
        return np.nan, (f"source label '{s}' is a RANGE of times, not a single "
                        f"time; time_h left blank")
    m = re.search(r"(\d+(?:\.\d+)?)\s*week", low)
    if m:
        return float(m.group(1)) * 7 * 24.0, f"source gave weeks ('{s}')"
    m = re.search(r"(\d+(?:\.\d+)?)\s*day", low)
    if m:
        return float(m.group(1)) * 24.0, f"source gave days ('{s}')"
    return np.nan, f"could not resolve a time from the source label '{s}'"


def read(d: Path) -> pd.DataFrame:
    wb = load_workbook(d / BOOK, data_only=True)
    rows: list[dict] = []

    def add(**kw):
        rows.append({"source_file": BOOK, "organism": "", "strain": "", **kw})

    # ---- 1a&b and 1c: groups across row 4, timepoints down column B -------
    for sheet, grp_row, first_row in (("1a&b", 4, 5), ("1c", 4, 5)):
        ws = wb[sheet]
        g = _filled(ws)
        groups = {c: _txt(g.get((grp_row, c), ""))
                  for c in range(3, ws.max_column + 1)}
        seen: dict[str, int] = {}
        for r in range(first_row, ws.max_row + 1):
            lbl = _txt(g.get((r, 2), ""))
            if not lbl or lbl.lower().startswith("sex"):
                continue
            t, why = parse_time(lbl)
            seen.clear()
            for c in range(3, ws.max_column + 1):
                val = _num(ws.cell(row=r, column=c).value)
                arm = groups.get(c, "")
                if np.isnan(val) or not arm:
                    continue
                seen[arm] = seen.get(arm, 0) + 1
                drug, dose, unit = parse_arm(arm)
                add(sheet=sheet, drug=drug, concentration=dose, conc_unit=unit,
                    arm=(f"{arm} | {lbl} [Fig {sheet}]" if np.isnan(t)
                         else f"{arm} [Fig {sheet}]"),
                    replicate=f"mouse{seen[arm]}", time_h=t, cfu_per_ml=val,
                    readout=LUNG,
                    notes=f"Figure {sheet}; {why}; dose is an administered "
                          f"mg/kg body-weight dose, not a medium concentration; "
                          f"the workbook does not name the organism")

    # ---- 3a: one row per condition, sex given per column ------------------
    ws = wb["3a"]
    g = _filled(ws)
    sexes = {c: _txt(g.get((3, c), "")) for c in range(3, ws.max_column + 1)}
    for r in range(4, ws.max_row + 1):
        lbl = _txt(g.get((r, 2), ""))
        if not lbl:
            continue
        t, why = parse_time(lbl)
        drug, dose, unit = parse_arm(lbl)
        n = 0
        for c in range(3, ws.max_column + 1):
            val = _num(ws.cell(row=r, column=c).value)
            if np.isnan(val):
                continue
            n += 1
            add(sheet="3a", drug=drug, concentration=dose, conc_unit=unit,
                arm=f"{lbl} [Fig 3a]", replicate=f"mouse{n}", time_h=t,
                cfu_per_ml=val,
                readout=LUNG,
                notes=f"Figure 3a; {why}; sex of this mouse as the sheet "
                      f"states it: {sexes.get(c, '') or 'not given'}; the "
                      f"workbook does not name the organism")

    # ---- 4a and 4b: therapy groups on row 3, sex row 4, times down B ------
    for sheet, readout in (("4a", LUNG), ("4b", MBLA)):
        ws = wb[sheet]
        g = _filled(ws)
        groups = {c: _txt(g.get((3, c), "")) for c in range(3, ws.max_column + 1)}
        sexes = {c: _txt(g.get((4, c), "")) for c in range(3, ws.max_column + 1)}
        for r in range(5, ws.max_row + 1):
            lbl = _txt(g.get((r, 2), ""))
            if not lbl:
                continue
            t, why = parse_time(lbl)
            seen: dict[str, int] = {}
            for c in range(3, ws.max_column + 1):
                val = _num(ws.cell(row=r, column=c).value)
                arm = groups.get(c, "")
                if np.isnan(val) or not arm:
                    continue
                seen[arm] = seen.get(arm, 0) + 1
                drug, dose, unit = parse_arm(arm)
                add(sheet=sheet, drug=drug, concentration=dose, conc_unit=unit,
                    arm=f"{arm} | {lbl} [Fig {sheet}]",
                    replicate=f"mouse{seen[arm]}",
                    time_h=t, cfu_per_ml=val, readout=readout,
                    notes=f"Figure {sheet}; {why}; sex of this mouse as the "
                          f"sheet states it: {sexes.get(c, '') or 'not given'}; "
                          f"the workbook does not name the organism")

    df = pd.DataFrame(rows)
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = np.nan if c in ("concentration", "time_h", "colonies",
                                    "dilution", "plated_volume_ul",
                                    "cfu_per_ml", "log10_cfu_per_ml",
                                    "floor_cfu_per_ml") else ""
    df["floor_basis"] = ("no plated volume, dilution or limit of detection "
                         "appears anywhere in this workbook")
    return df[COLUMNS]
