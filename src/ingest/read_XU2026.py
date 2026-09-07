"""XU2026 -- Communications Biology supplementary data (42003_2026_9947_MOESM3).

Questiomycin A (QM) against Mycobacterium tuberculosis. One workbook, one sheet
per figure, several small blocks per sheet. The CFU blocks are read; the
RT-qPCR, fluorescence, ATP-luminescence and plasma-PK blocks are not (they are
listed in NOT_READ below).

WHAT THE DEPOSIT DOES AND DOES NOT SAY
  - It says "CFU counts" everywhere and never "CFU/mL". The counts are
    therefore returned in cfu_per_ml with readout recording that the source did
    not state the denominator. No volume, no dilution and no plating detail
    appear anywhere in the workbook, so floor_cfu_per_ml stays blank.
  - Many entries are a literal 0. The workbook gives no detection limit, so a
    zero is returned as zero and left uncensored -- that ambiguity is the
    measurement, not a defect to be patched.
  - Time is in days on fig.1, fig.2, fig.6 and the supplementary sheet, and in
    hours on 'table 2'. Everything is converted to hours.
  - The organism is recorded EXACTLY as each block's own title writes it, and
    the workbook is not consistent: fig.2 spells out "Mycobacterium
    tuberculosis H37Rv", 'table 2' and fig.6b write only "Mtb" (fig.6a, an
    RT-qPCR block that is not read, writes "M. tuberculosis H37Rv"), and the
    supplementary sheet writes "Mtb". fig.1 names no species at all -- only the
    strain -- so organism is blank on every fig.1 row. Nothing is carried from
    one sheet to another.
  - Supplementary Figure S3 lays TWO stress blocks side by side on one header
    row (columns B..H and L..R). The column walk is bounded to each block's own
    columns; before that bound the left block's header carried on into the right
    block's columns and read the right block's numbers a second time under the
    left block's condition label.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.ingest import COLUMNS

BOOK = "_unpacked/PMC13201773_supplementary/42003_2026_9947_MOESM3_ESM.xlsx"

NOT_READ = (
    "fig.2b (intracellular CFU) has two unlabelled value columns per replicate "
    "and the sheet does not say what distinguishes them; fig.3, fig.4, fig.6a/c/f, "
    "Fig. S6, Fig. S7 and 'table 3' are RT-qPCR, fluorescence, ATP luminescence "
    "or plasma PK, not bacterial counts"
)

READOUT = "CFU (source says 'CFU counts'; it never states the denominator)"

_CTRL = {"ctrl", "control", "growth control", "dmso", "vehicle"}


def _txt(v) -> str:
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def _num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return np.nan


def parse_label(label: str) -> tuple[str, str, float, str]:
    """(strain, drug, concentration, conc_unit) from a column or row label.

    Only what the label itself contains. A label the parser cannot decompose
    comes back with blank drug and NaN concentration, and the caller keeps the
    verbatim label in `arm`.
    """
    s = label.replace("\xa0", " ").replace("×", "x").replace("μ", "u")
    s = s.replace("µ", "u").strip()
    if s.lower().strip() in _CTRL:
        return "", "", np.nan, ""

    strain = ""
    m = re.match(r"^\s*([A-Za-z0-9_\-]+)\s*\((.+)\)\s*$", s)
    if m and re.search(r"MIC|g/mL|/mL", m.group(2)):
        head, inner = m.group(1), m.group(2)
        # 'QM(8 x MIC)' -> head is the drug; 'pMV261 (QM 1 x MIC)' -> head is a strain
        if re.match(r"^\s*[\d./]+\s*x\s*MIC\s*$", inner, re.I) or \
                re.match(r"^\s*[\d.]+\s*ug/mL\s*$", inner, re.I):
            s = f"{head} {inner}"
        else:
            strain, s = head, inner

    drug = ""
    for name in ("QM", "RFP", "BDQ", "INH", "Rifampin", "rifampicin", "DCCD"):
        if re.search(rf"\b{name}\b", s, re.I):
            drug = name
            break

    conc, unit = np.nan, ""
    m = re.search(r"(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)\s*x\s*MIC", s, re.I)
    if m:
        conc, unit = float(m.group(1)) / float(m.group(2)), "xMIC"
    else:
        m = re.search(r"(\d+(?:\.\d+)?)\s*x\s*MIC", s, re.I)
        if m:
            conc, unit = float(m.group(1)), "xMIC"
        else:
            m = re.search(r"(\d+(?:\.\d+)?)\s*(ug/mL|uM|mg/mL)", s, re.I)
            if m:
                conc, unit = float(m.group(1)), m.group(2).replace("ug", "ug")
    if not strain and not drug and not unit:
        strain = s
    return strain, drug, conc, unit


def _organism(title: str) -> str:
    """The species EXACTLY as this block's own title writes it, or blank.

    The workbook is inconsistent: fig.2 spells 'Mycobacterium tuberculosis'
    out, fig.6b and 'table 2' write only 'Mtb', fig.1 names no species at all.
    Each block gets what its own title says; nothing is carried between sheets.
    """
    for form in ("Mycobacterium tuberculosis", "M. tuberculosis", "Mtb"):
        if form in title:
            return form
    return ""


def _strain(title: str) -> str:
    """The strain the block's own title names, or blank."""
    return "H37Rv" if "H37Rv" in title else ""


def _groups(ws, row: int, first_col: int = 3, last_col: int | None = None):
    """Column -> group label, spans carried forward across merged headers.

    `last_col` bounds the walk. It matters on the Supplementary Figure sheet,
    which puts TWO blocks side by side on the same header row: without a bound
    the left block's header carry-forward runs on into the right block's columns
    and the right block's numbers are read a second time under the left block's
    condition label.
    """
    out, cur = {}, None
    for col in range(first_col, (last_col or ws.max_column) + 1):
        v = _txt(ws.cell(row=row, column=col).value)
        if v:
            cur = v
        out[col] = cur
    return out


def _day_blocks(ws, key: str = "Days"):
    """Yield (title, header_row) for every block whose column B says `key`."""
    for row in range(1, ws.max_row + 1):
        if _txt(ws.cell(row=row, column=2).value) == key:
            title, fig = "", ""
            for up in range(row - 1, 0, -1):
                b = _txt(ws.cell(row=up, column=2).value)
                if b:
                    title = b
                    fig = _txt(ws.cell(row=up, column=1).value)
                    break
            yield title, fig, row


def _read_day_block(ws, title, fig, hrow, sheet, src, organism, block_drug,
                    time_unit, rows, block_tag=""):
    groups = _groups(ws, hrow)
    reps = {c: _txt(ws.cell(row=hrow + 1, column=c).value)
            for c in range(3, ws.max_column + 1)}
    r = hrow + 2
    while r <= ws.max_row:
        t = _num(ws.cell(row=r, column=2).value)
        if np.isnan(t):
            break
        for col, arm in groups.items():
            if not arm:
                continue
            val = _num(ws.cell(row=r, column=col).value)
            if np.isnan(val):
                continue
            strain, drug, conc, unit = parse_label(arm)
            is_ctrl = arm.replace("\xa0", " ").strip().lower() in _CTRL
            rows.append({
                "source_file": src, "sheet": sheet,
                "organism": organism, "strain": strain,
                "drug": "" if is_ctrl else (drug or block_drug),
                "concentration": conc,
                "conc_unit": unit,
                "arm": f"{arm} [{block_tag}]" if block_tag else arm,
                "replicate": reps.get(col, ""),
                "time_h": t * (24.0 if time_unit == "day" else 1.0),
                "cfu_per_ml": val, "readout": READOUT,
                "notes": f"{fig + ': ' if fig else ''}{title}"
                         f" | time given in {'days' if time_unit == 'day' else 'hours'}"
                         f" by the source",
            })
        r += 1


def read(d: Path) -> pd.DataFrame:
    path = d / BOOK
    wb = load_workbook(path, data_only=True)
    src = BOOK
    rows: list[dict] = []

    # ---- fig.1: five time-kill blocks, drug named in the block title -------
    ws = wb["fig.1"]
    for title, fig, hrow in _day_blocks(ws):
        drug = "QM" if re.search(r"\bQM\b", title) else (
            "RFP" if re.search(r"\bRFP\b", title) else "")
        strain = title.split("–")[0].split("-")[0].strip() if "–" in title else ""
        if "Pre-XDR" in title:
            strain = title.split("–")[0].strip()
        n0 = len(rows)
        _read_day_block(ws, title, fig, hrow, "fig.1", src, "", drug, "day", rows,
                        block_tag=f"{strain} {drug}".strip())
        for rec in rows[n0:]:
            if not rec["strain"] or rec["strain"] in _CTRL:
                rec["strain"] = strain
            rec["notes"] += ("; the fig.1 sheet names the strain but not the "
                             "species -- other sheets of this workbook say "
                             "Mycobacterium tuberculosis, so organism is left blank here")

    # ---- fig.2: three pH blocks (7.0, 6.0, 5.0) ---------------------------
    ws = wb["fig.2"]
    for title, fig, hrow in _day_blocks(ws):
        org, st = _organism(title), _strain(title)
        ph = re.search(r"pH\s*[\d.]+", title)
        n0 = len(rows)
        _read_day_block(ws, title, fig, hrow, "fig.2", src, org, "", "day", rows,
                        block_tag=ph.group(0).rstrip(".") if ph else (fig or ""))
        for rec in rows[n0:]:
            rec["strain"] = rec["strain"] or st

    # ---- fig.2c: nutrient-starved non-replicating Mtb, no time stated ------
    ws = wb["fig.2"]
    hdr = None
    for row in range(1, ws.max_row + 1):
        if _txt(ws.cell(row=row, column=3).value) == "DMSO" and \
                _txt(ws.cell(row=row + 1, column=2).value).startswith("Replicate"):
            hdr = row
            break
    if hdr:
        title = ""
        for up in range(hdr - 1, 0, -1):
            b = _txt(ws.cell(row=up, column=2).value)
            if b:
                title = b
                break
        cols = {c: _txt(ws.cell(row=hdr, column=c).value)
                for c in range(3, ws.max_column + 1)
                if _txt(ws.cell(row=hdr, column=c).value)}
        r = hdr + 1
        while _txt(ws.cell(row=r, column=2).value).startswith("Replicate"):
            rep = _txt(ws.cell(row=r, column=2).value)
            for col, arm in cols.items():
                val = _num(ws.cell(row=r, column=col).value)
                if np.isnan(val):
                    continue
                strain, drug, conc, unit = parse_label(arm)
                rows.append({
                    "source_file": src, "sheet": "fig.2",
                    "organism": _organism(title),
                    "strain": "", "drug": drug, "concentration": conc,
                    "conc_unit": unit,
                    "arm": f"{arm} [NR-Mtb]", "replicate": rep,
                    "time_h": np.nan, "cfu_per_ml": val, "readout": READOUT,
                    "notes": f"Fig.2c: {title} | the sheet gives no exposure "
                             f"time for this block, so time_h is blank",
                })
            r += 1

    # ---- fig.6b: FabD over-expression, days 0-7 ---------------------------
    ws = wb["fig.6"]
    for title, fig, hrow in _day_blocks(ws):
        _read_day_block(ws, title, fig, hrow, "fig.6", src,
                        _organism(title), "", "day", rows,
                        block_tag=fig or "Fig.6b")

    # ---- table 2: post-antibiotic effect, time already in hours -----------
    ws = wb["table 2"]
    for title, fig, hrow in _day_blocks(ws, key="Time (h)"):
        n0 = len(rows)
        _read_day_block(ws, title, fig, hrow, "table 2", src,
                        _organism(title), "", "hour", rows,
                        block_tag="PAE, 2 h exposure")
        st = _strain(title)
        for rec in rows[n0:]:
            rec["strain"] = rec["strain"] or st

    # ---- Supplementary Figure S3: four stress blocks, Day 1 / Day 7 -------
    ws = wb["Supplementary Figure"]
    LEADS = (2, 12)                                # left block B.., right block L..
    for hrow in range(1, ws.max_row + 1):
        for i, lead in enumerate(LEADS):
            head_col = lead + 1
            last_col = LEADS[i + 1] - 1 if i + 1 < len(LEADS) else ws.max_column
            if not _txt(ws.cell(row=hrow, column=head_col).value).startswith("Day "):
                continue
            if not _txt(ws.cell(row=hrow + 1, column=head_col).value).startswith("Replicate"):
                continue
            title = ""
            for up in range(hrow - 1, 0, -1):
                b = _txt(ws.cell(row=up, column=lead).value)
                if b:
                    title = b
                    break
            days = _groups(ws, hrow, first_col=head_col, last_col=last_col)
            reps = {c: _txt(ws.cell(row=hrow + 1, column=c).value)
                    for c in range(head_col, last_col + 1)}
            r = hrow + 2
            while True:
                arm = _txt(ws.cell(row=r, column=lead).value)
                if not arm:
                    break
                for col, day in days.items():
                    if not day or not day.startswith("Day "):
                        continue
                    val = _num(ws.cell(row=r, column=col).value)
                    if np.isnan(val):
                        continue
                    strain, drug, conc, unit = parse_label(arm)
                    rows.append({
                        "source_file": src, "sheet": "Supplementary Figure",
                        "organism": "Mtb", "strain": "",
                        "drug": drug, "concentration": conc, "conc_unit": unit,
                        "arm": f"{arm} [{title}]", "replicate": reps.get(col, ""),
                        "time_h": float(day.split()[1]) * 24.0,
                        "cfu_per_ml": val, "readout": READOUT,
                        "notes": f"Fig. S3: {title} | organism as the sheet "
                                 f"writes it ('against Mtb'); time given in days",
                    })
                r += 1

    df = pd.DataFrame(rows)
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = np.nan if c in ("concentration", "time_h", "colonies",
                                    "dilution", "plated_volume_ul",
                                    "cfu_per_ml", "log10_cfu_per_ml",
                                    "floor_cfu_per_ml") else ""
    df["floor_basis"] = ("no plated volume, dilution or limit of detection "
                         "appears anywhere in this workbook")
    return df[COLUMNS]
