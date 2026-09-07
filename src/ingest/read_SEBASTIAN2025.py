"""SEBASTIAN2025 -- Nature Communications Source Data (41467_2025_67152_MOESM9).

A transcription-factor induction (TFI) / CRISPRi-knockdown library in
Mycobacterium tuberculosis, killed with rifampicin and five other drugs over
days. One sheet per figure; most sheets carry several small blocks. Only the
blocks that are colony counts are read.

WHAT THE DEPOSIT DOES AND DOES NOT SAY
  - THE WORKBOOK IS NOT UNIFORM ABOUT WHAT IT MEASURED, and readout records
    exactly what each block's own sheet says, no more:
      'Figure 1 B'          "CFU/mL"                   -> per mL, stated
      Fig. S8G              "DAY0 CFU/mL"              -> per mL, stated
      Figure 2N, Fig. S5E/F "Estiated CFU per well"    -> per WELL, stated,
                                                         NOT per mL
      Figure 3G-I           "CFU  after 24 hours of H2O2 treatment" in each
                            block title -> CFU, no denominator
      Figure 2A-I           the sheet does not head its value columns, but M2
                            writes the derived quantity as "ATc(Day 4 CFU/ Day
                            0 CFU) / No-ATc(...)", which names those same
                            columns CFU -> CFU, no denominator
      Figure 4, Figure 6    THE SHEET NEVER SAYS WHAT WAS MEASURED. The word
                            CFU does not occur anywhere on either sheet. The
                            blocks have a day axis, a construct label and
                            (Figure 4) a drug, and the magnitudes look like the
                            CFU blocks elsewhere in the workbook -- but the
                            deposit does not say so, so readout records that
                            the source did not state it rather than asserting
                            CFU. The rows are kept, because everything else
                            about them IS stated; anyone who needs counts must
                            filter on readout.
    Nothing is converted between denominators anywhere.
  - No sheet anywhere in the workbook names the organism. Gene identifiers are
    Rv numbers and the constructs are mycobacterial, but the file does not say
    it, so organism is blank on every row.
  - No plated volume, dilution, or limit of detection appears anywhere.
    floor_cfu_per_ml is blank throughout. Zeros are returned as zeros.
  - Time is in days everywhere except the H2O2 blocks of Figure 3G-I, where the
    sheet says "CFU after 24 hours of H2O2 treatment". Converted to hours.
  - 'Figure 1 B' contains the literal string 'RifR' in place of a count in five
    cells. Those cells are skipped; the count is on the returned frame as
    df.attrs["skipped_rifr"].
  - 'Figure 4' has FIVE cells reading exactly 40000 at day 12 in its two INH
    blocks (4C: E20, F20; 4I: L20, O20, P20), where every sibling in the same
    block is of order 10^2.
    The deposit does not explain them. They are returned as deposited, with the
    anomaly written into `notes` on those rows.
  - 'Figure 4' repeats ONE day-0 measurement across all six of a construct's
    drug blocks: rows 4/11/18/25/32/39 are byte-identical for PLJR_cydA and
    likewise for PLJR_icl1. The 72 day-0 rows therefore carry only 12 distinct
    numbers. They are returned (each block needs its own t = 0) with the
    repetition spelled out in notes so they can be de-duplicated before
    pooling.
  - 'Figure 1 B' gives no replicate labels and a different number of values per
    day (4, 4, 8, 12), so its replicate index is only a column position and the
    columns are NOT paired across days. `notes` says so on every such row.

NOT READ, and why -- see NOT_READ.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.ingest import COLUMNS

BOOK = "_unpacked/PMC12800208_supplementary/41467_2025_67152_MOESM9_ESM.xlsx"

NOT_READ = (
    "Figure 1C/1D/2J/2K/2L/2M/5H/5I/5J-M, Supplementary Figures 1E, 4, 5A-D, "
    "7A-C, 10 are MiSeq read counts, frequencies or normalised fractions, not "
    "colony counts; Figure 2A-I columns M-Y are Day4/Day0 ratios derived from "
    "the same block and are not re-read; Figure 3A-F, 5A, Supplementary Figures "
    "2, 3, 6, 7D-E, 9 are CellROX/flow/RT-PCR/MIC readouts; Supplementary Figure "
    "8A-C (a menadione dose series) and 8D (a TFI H2O2 panel) are blocks "
    "of bare numbers -- the sheet gives them no units, no readout label and no "
    "exposure time, and the word CFU appears on that sheet only over the S8G "
    "block -- so what they measure cannot be established from the deposit and "
    "they are left out rather than assumed to be counts; Supplementary Figure "
    "8E/8F are CellROX fluorescence."
)

NO_ORG = ("no sheet in this workbook names the organism, so organism is blank")

UNLABELLED = ("not stated by the source -- this sheet never says what was "
              "measured and the word CFU does not appear anywhere on it")

SHARED_DAY0 = ("; NOTE: this day-0 row is identical in all SIX drug blocks of "
               "this construct (4A-4F for PLJR_cydA, 4G-4L for PLJR_icl1) -- "
               "the deposit repeats one shared day-0 measurement, so these are "
               "six copies of three numbers, not eighteen independent "
               "readings; de-duplicate before pooling")


def _txt(v) -> str:
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def _num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return np.nan


def _day(v) -> float:
    """Days -> hours from a cell that may be a number or 'Day 4'."""
    n = _num(v)
    if not np.isnan(n):
        return n * 24.0
    m = re.search(r"(-?\d+(?:\.\d+)?)", _txt(v))
    return float(m.group(1)) * 24.0 if m else np.nan


def _spans(ws, row: int, cols: range) -> dict[int, str]:
    out, cur = {}, ""
    for c in cols:
        v = _txt(ws.cell(row=row, column=c).value)
        if v:
            cur = v
        out[c] = cur
    return out


def _strain_arm(label: str) -> tuple[str, str]:
    """(strain, induced-or-not) from labels like 'sigI-ATc' or 'cydA_3_ATc'."""
    s = _txt(label)
    base = re.sub(r"[-_ ]?(ATc|Atc)\b", "", s).strip(" -_")
    return base, s


def read(d: Path) -> pd.DataFrame:
    wb = load_workbook(d / BOOK, data_only=True)
    rows: list[dict] = []
    skipped_rifr = 0

    def add(**kw):
        rec = {"source_file": BOOK, "organism": "", **kw}
        rows.append(rec)

    # ------------------------------------------------ Figure 1 B ----------
    ws = wb["Figure 1 B"]
    title = _txt(ws["C1"].value)
    blocks = {"Control (CFU/mL)": range(3, 15), "ATc Induced (CFU/mL)": range(15, 27)}
    for grp, cols in blocks.items():
        for r in range(3, ws.max_row + 1):
            lbl = _txt(ws.cell(row=r, column=2).value)
            if not lbl.lower().startswith("day"):
                continue
            for i, c in enumerate(cols, start=1):
                raw = ws.cell(row=r, column=c).value
                if raw is None:
                    continue
                val = _num(raw)
                if np.isnan(val):
                    skipped_rifr += 1
                    continue
                add(sheet="Figure 1 B", strain="TRIP library", drug="Rifampicin",
                    arm=f"{grp} [{title.strip()}]", replicate=f"col{i}",
                    time_h=_day(lbl), cfu_per_ml=val, readout="CFU/mL",
                    notes=f"{title.strip()} | the sheet gives no replicate "
                          f"labels and the number of values per day differs "
                          f"(4, 4, 8, 12), so the replicate index is only the "
                          f"column position and columns are NOT paired across "
                          f"days; cells reading 'RifR' were skipped; " + NO_ORG)

    # ------------------------------------ Figure 2A-I (blocks 2A .. 2H) ---
    ws = wb["Figure 2A-I"]
    for r in range(1, ws.max_row + 1):
        if _txt(ws.cell(row=r, column=2).value) != "Day of Rif treatment":
            continue
        groups = _spans(ws, r, range(3, 11))       # C..J only; M.. are ratios
        rr = r + 1
        while rr <= ws.max_row:
            t = _day(ws.cell(row=rr, column=2).value)
            if np.isnan(t):
                break
            for c, grp in groups.items():
                if not grp:
                    continue
                val = _num(ws.cell(row=rr, column=c).value)
                if np.isnan(val):
                    continue
                strain, arm = _strain_arm(grp)
                base = 3 if c < 7 else 7
                add(sheet="Figure 2A-I", strain=strain, drug="Rif",
                    arm=arm, replicate=f"rep{c - base + 1}", time_h=t,
                    cfu_per_ml=val,
                    readout="CFU (the sheet does not head its value columns, "
                            "but M2 defines the derived ratio as 'ATc(Day 4 "
                            "CFU/ Day 0 CFU) / No-ATc(...)', which names these "
                            "columns CFU; no denominator is stated)",
                    notes="header reads 'Day of Rif treatment'; days converted "
                          "to hours; " + NO_ORG)
            rr += 1

    # ------------------- Figure 2N and Supplementary Figure 5E,F ----------
    for sheet in ("Figure 2N", "Supplementary Figure 5E, F"):
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            if _txt(ws.cell(row=r, column=3).value) != "Day":
                continue
            groups = _spans(ws, r, range(4, 10))
            rr = r + 1
            while rr <= ws.max_row:
                t = _day(ws.cell(row=rr, column=3).value)
                if np.isnan(t):
                    break
                for c, grp in groups.items():
                    val = _num(ws.cell(row=rr, column=c).value)
                    if not grp or np.isnan(val):
                        continue
                    strain, arm = _strain_arm(grp)
                    base = 4 if c < 7 else 7
                    add(sheet=sheet, strain=strain, drug="", arm=arm,
                        replicate=f"rep{c - base + 1}", time_h=t,
                        cfu_per_ml=val,
                        readout="estimated CFU per well (the sheet's own "
                                "wording; NOT per mL)",
                        notes="the sheet heads this block 'Estiated CFU per "
                              "well' and names no drug, so drug is blank; " + NO_ORG)
                rr += 1

    # ------------------------------------------------ Figure 4 ------------
    ws = wb["Figure 4"]
    for r in range(1, ws.max_row + 1):
        for lead, tcol in ((2, 3), (11, 12)):      # B|C.. and K|L..
            if _txt(ws.cell(row=r, column=lead).value) != "Day of treatment":
                continue
            title = _txt(ws.cell(row=r - 1, column=tcol).value)
            groups = _spans(ws, r, range(tcol, tcol + 6))
            rr = r + 1
            while rr <= ws.max_row:
                t = _day(ws.cell(row=rr, column=lead).value)
                if np.isnan(t):
                    break
                for c, grp in groups.items():
                    val = _num(ws.cell(row=rr, column=c).value)
                    if not grp or np.isnan(val):
                        continue
                    m = re.match(r"Figure\s*\d+[A-Z]_(.+?)_([A-Za-z0-9]+)$", title)
                    strain = m.group(1) if m else ""
                    drug = m.group(2) if m else ""
                    base = tcol if c < tcol + 3 else tcol + 3
                    shared0 = SHARED_DAY0 if t == 0 else ""
                    odd = ("; CAUTION: this cell reads exactly 40000 at day 12 "
                           "where its siblings in the same block are of order "
                           "10^2 -- five such cells occur in the two INH blocks "
                           "and the deposit does not explain them; returned as "
                           "deposited" if val == 40000 else "")
                    add(sheet="Figure 4", strain=strain, drug=drug,
                        arm=f"{grp} [{title}]", replicate=f"rep{c - base + 1}",
                        time_h=t, cfu_per_ml=val,
                        readout=UNLABELLED,
                        notes=f"{title}; drug and construct taken from that "
                              f"block title; days converted to hours; "
                              + NO_ORG + shared0 + odd)
                rr += 1

    # ------------------------------------------------ Figure 6 ------------
    ws = wb["Figure 6"]
    def _is_day(row, col):
        return _txt(ws.cell(row=row, column=col).value).lower().startswith("day")

    for lead in (3, 11):                            # day label in C or K
        r = 1
        while r <= ws.max_row:
            if not _is_day(r, lead) or (r > 1 and _is_day(r - 1, lead)):
                r += 1
                continue
            hdr = r - 1
            groups = _spans(ws, hdr, range(lead + 1, lead + 7))
            if any(g in ("I/UI", "avg.") for g in groups.values()):
                r += 1
                continue
            while r <= ws.max_row and _is_day(r, lead):
                t = _day(ws.cell(row=r, column=lead).value)
                for c, grp in groups.items():
                    val = _num(ws.cell(row=r, column=c).value)
                    if not grp or np.isnan(val):
                        continue
                    strain, arm = _strain_arm(grp)
                    base = lead + 1 if c < lead + 4 else lead + 4
                    add(sheet="Figure 6", strain=strain, drug="", arm=arm,
                        replicate=f"rep{c - base + 1}", time_h=t,
                        cfu_per_ml=val,
                        readout=UNLABELLED,
                        notes="Figure 6 gives days and construct labels but "
                              "names no drug, so drug is blank; " + NO_ORG)
                r += 1

    # ---------------------------------------------- Figure 3G-I -----------
    ws = wb["Figure 3G-I"]
    for r in range(1, ws.max_row + 1):
        if _txt(ws.cell(row=r, column=2).value) != "H2O2 (%)":
            continue
        title = _txt(ws.cell(row=r - 1, column=3).value)
        groups = _spans(ws, r, range(3, 9))
        rr = r + 1
        while rr <= ws.max_row:
            conc = _num(ws.cell(row=rr, column=2).value)
            if np.isnan(conc):
                break
            for c, grp in groups.items():
                val = _num(ws.cell(row=rr, column=c).value)
                if not grp or np.isnan(val):
                    continue
                base = 3 if c < 6 else 6
                add(sheet="Figure 3G-I",
                    strain=title.split(" knockdown")[0].split(" (")[0].strip(),
                    drug="H2O2", concentration=conc, conc_unit="%",
                    arm=f"{grp} [{title}]", replicate=f"rep{c - base + 1}",
                    time_h=24.0, cfu_per_ml=val,
                    readout="CFU (the block title says 'CFU  after 24 hours of "
                            "H2O2 treatment'; no denominator is stated)",
                    notes=f"{title} -- the sheet states 24 hours of H2O2 "
                          f"treatment, which is the only time it gives; " + NO_ORG)
            rr += 1

    # --------------------------------- Supplementary Figure 8, panel G ----
    ws = wb["Supplementary Figure 8"]
    anchor = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if _txt(ws.cell(row=r, column=c).value).upper().startswith("DAY0 CFU/ML"):
                anchor = (r, c)
    if anchor:
        hrow, c0 = anchor
        day1 = _txt(ws.cell(row=hrow, column=c0 + 3).value)
        conc = 0.01 if "0.01%" in day1 else np.nan
        rr = hrow + 1
        while rr <= ws.max_row:
            lbl = _txt(ws.cell(row=rr, column=c0 - 1).value)
            if not lbl:
                break
            strain, arm = _strain_arm(lbl)
            for off, (t, drug) in enumerate(((0.0, ""), (24.0, "H2O2"))):
                for i in range(3):
                    val = _num(ws.cell(row=rr, column=c0 + off * 3 + i).value)
                    if np.isnan(val):
                        continue
                    add(sheet="Supplementary Figure 8", strain=strain,
                        drug=drug,
                        concentration=conc if drug else np.nan,
                        conc_unit="%" if drug else "",
                        arm=f"{arm} [Fig. S8G]", replicate=f"rep{i + 1}",
                        time_h=t, cfu_per_ml=val, readout="CFU/mL",
                        notes=f"Fig. S8G: columns headed '{_txt(ws.cell(row=hrow, column=c0).value)}'"
                              f" and '{day1}'; " + NO_ORG)
            rr += 1

    df = pd.DataFrame(rows)
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = np.nan if c in ("concentration", "time_h", "colonies",
                                    "dilution", "plated_volume_ul",
                                    "cfu_per_ml", "log10_cfu_per_ml",
                                    "floor_cfu_per_ml") else ""
    df["floor_basis"] = ("no plated volume, dilution or limit of detection "
                         "appears anywhere in this workbook")
    df.attrs["skipped_rifr"] = skipped_rifr
    return df[COLUMNS]
