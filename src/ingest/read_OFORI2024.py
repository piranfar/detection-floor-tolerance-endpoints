"""OFORI2024 -- Nature Communications Source Data (41467_2024_53933_MOESM8).

katG and bedaquiline killing in Mycobacterium tuberculosis. The Source Data
workbook is almost entirely dose-response curves and expression values; exactly
one block in it is a colony count with a time axis, and that is all this reader
returns.

WHAT IS READ
  Sheet 'Figure 1c', the block headed "CFU/mL" with a "Day" column: ten strains
  (H37Rv and nine TDR-* clinical isolates) at day 0 and day 30, one value per
  strain per day. That is 20 readings.

WHAT IS DELIBERATELY NOT READ, AND WHY
  - 'Figure 1c' also carries a derived "30-Day % Survival" block computed from
    the same twenty numbers. Reading it would double-count.
  - 'Figure 1e' and 'Figure 1f' have a "Day" column (0, 3, 8, 16, 30) and three
    values per strain, but the sheet states NO unit and no readout for those
    values. They are near 100 at day 0 and fall to 0.003, which is consistent
    with a percentage of the day-0 count, but the deposit does not say so. They
    are left out rather than guessed at.
  - Figures 1a, 1b, 1d, 2a, 2b, 2d, 3b, 3c, 4b, 4c, 4d, 5b, 5c and Supplementary
    Figures 1a, 1b, 2a, 2b, 2c, 3c, 4b, 6c, 6d are drug- or H2O2-concentration
    series with no time axis at all, and the sheets do not say what the values
    are (the small ones sit at ~0.045 and the BacTiter-Glo ones are labelled as
    luminescence). Without a time and without a stated readout they are not
    readings in this schema's sense.
  - Figures 2c, 2e, 2f, 2g, 3a, 4a, 5a, Supplementary Figures 3a, 3b, 5a-c, 6a,
    6b are expression, CellROX fluorescence or statistics.
  - MOESM4 (normalised expression), MOESM5 (RNAseq clusters, GO) and MOESM6
    (iEK1011 flux simulations) contain no colony counts.

WHAT THE DEPOSIT DOES NOT SAY
  - The block gives "CFU/mL" and "Day" and nothing else: no drug is named
    anywhere on that sheet, so `drug` is blank on every row. (The statistics
    block above it compares INH-susceptible with INH-resistant isolates, which
    describes the strains, not the treatment.)
  - No organism is named. Strain labels are kept as given.
  - No plated volume, dilution or limit of detection anywhere, so
    floor_cfu_per_ml is blank.
  - There is one value per strain per day: no replicates.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.ingest import COLUMNS

BOOK = "_unpacked/PMC11561320_supplementary/41467_2024_53933_MOESM8_ESM.xlsx"

NOT_READ = (
    "Figure 1e/1f have a day axis but the sheet states no unit or readout for "
    "their values; every other sheet is a concentration series with no time, an "
    "expression or fluorescence measurement, or a statistics block; the derived "
    "'30-Day % Survival' block of Figure 1c is not re-read"
)


def _txt(v) -> str:
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def _num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return np.nan


def read(d: Path) -> pd.DataFrame:
    wb = load_workbook(d / BOOK, data_only=True)
    ws = wb["Figure 1c"]
    rows: list[dict] = []

    unit_row = next((r for r in range(1, ws.max_row + 1)
                     if _txt(ws.cell(row=r, column=1).value) == "CFU/mL"), None)
    hdr = next((r for r in range(1, ws.max_row + 1)
                if _txt(ws.cell(row=r, column=1).value) == "Day"), None)
    if unit_row is None or hdr is None or hdr != unit_row + 1:
        return pd.DataFrame(columns=COLUMNS)

    strains = {c: _txt(ws.cell(row=hdr, column=c).value)
               for c in range(2, ws.max_column + 1)
               if _txt(ws.cell(row=hdr, column=c).value)}
    r = hdr + 1
    while r <= ws.max_row:
        day = _num(ws.cell(row=r, column=1).value)
        if np.isnan(day):
            break
        for c, strain in strains.items():
            val = _num(ws.cell(row=r, column=c).value)
            if np.isnan(val):
                continue
            rows.append({
                "source_file": BOOK, "sheet": "Figure 1c",
                "organism": "", "strain": strain, "drug": "",
                "arm": strain, "replicate": "",
                "time_h": day * 24.0, "cfu_per_ml": val,
                "readout": "CFU/mL",
                "notes": "block headed 'CFU/mL' with a 'Day' column; days "
                         "converted to hours. The sheet names NO drug for this "
                         "killing experiment, so drug is blank; it names no "
                         "organism either. One value per strain per day -- the "
                         "deposit gives no replicates.",
            })
        r += 1

    df = pd.DataFrame(rows)
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = np.nan if c in ("concentration", "time_h", "colonies",
                                    "dilution", "plated_volume_ul",
                                    "cfu_per_ml", "log10_cfu_per_ml",
                                    "floor_cfu_per_ml") else ""
    df["floor_basis"] = ("no plated volume, dilution or limit of detection "
                         "appears anywhere in this workbook")
    return df[COLUMNS]
