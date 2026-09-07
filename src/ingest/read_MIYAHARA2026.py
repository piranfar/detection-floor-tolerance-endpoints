"""MIYAHARA2026 -- figshare 10.6084/m9.figshare.30780101, figure source data.

Twelve workbooks, one per figure or figure group. Exactly one sheet in the whole
deposit holds plate-count kill curves: sheet `FigS1` of Fig1_S1.xlsx, laid out
as three blocks keyed in column A --

    S1a   Time (min)   UPEC (11 columns)        MG1655 (12 columns)
    S1b   Time (min)   12.5 / 25 / 50 / 100 (5, 5, 4, 5 columns)
    S1c   <no header>  UPEC_YFP (3 columns)

one column per biological replicate, one row per timepoint.

Two further panels are returned as NON-count readings, because the sheet itself
says what its numbers are and gives them a time axis:

    Fig1b   `Time (min)` / `Cell count` -- a count of cells in a microscope
            field of view, NOT per mL. cfu_per_ml is blank on those rows and the
            number rides in notes; nothing turns a field count into a density.
    FigS4a  `Survival fraction`, six EHEC replicates, t = 0 written as 1. The
            denominator is nowhere in the deposit, so no count is reconstructed.

WHAT IS LEFT, and exactly what each sheet says on its face -- so that nobody has
to reopen twelve workbooks to find out. The line drawn here is: a panel is
returned only if the sheet states what its values ARE. These do not.

    Fig1c            `Kill rates` against `Time (min)`, plus three replicate
                     columns headed 08 / 09 / 10; a rate, not a reading
    Fig2_S2          ten `Persister Lineage` sheets: `Time`, `Cell size` and a
                     bare `AMP` column -- single-cell traces
    Fig2b/2c         `Growth rates (min-1)`, `Cell Size (um2) at birth`, split
                     Died / Survived; no time axis
    Fig3c            `Time(h)` against groups `M9` and `M9+40X aMG`, values
                     starting at 100. The sheet NEVER names the quantity: it is
                     not labelled per cent, survival, OD or anything else, so
                     there is no honest readout to give it
    Fig 3d           two bare columns `M9` / `M9+40X aMG`, no axis, no quantity
    Fig4a/4b         `Growth rates (min-1) in M9` / `in M9+40x aMG`
    Fig4c/4d         `Bin Center` histogram bins of those rates
    Fig4e/4f         `growth rate` against `Lysis rate`, per cell
    Fig5c/5d,
    Fig6c/6d         per-cell RpoS or TolC fluorescence and lysis rates
    Fig5e/5f,
    Fig6e/6f         per-cell traces against `Time` / `Time (min)`, ~200 columns
                     each: single cells, not cultures
    FigS3            one unlabelled 140 x 16 grid, no header row at all
    FigS4c/S4d       `*EHEC_AMP_Lysis count` and an unlabelled grid against
                     `Time (min)` -- microscope lysis counts
    FigS5a           `Time(h)` against `M9` / `M9+40X aMG`, values 0.025 rising
                     to ~1. The sheet does NOT say OD600 or anything else; the
                     quantity is unnamed, so it is not returned
    FigS5b/S5c       same shape as Fig3c, values starting at 100, quantity
                     unnamed
    FigS6a/S6b       `Growth rate (min-1)`, `Lysis rates (min-1)`
    FigS9b/S9c       `Fluorescence (RpoS levels)`, `Lysis rates (min-1)`

WHAT THE DEPOSIT DOES NOT SAY, and is therefore blank:

  * the drug. No cell in any of the twelve workbooks names an antibiotic in
    full. The closest are a bare column heading `AMP` in Fig2_S2 and
    `*EHEC_AMP_Lysis count` in FigS4c, neither on a sheet returned here as a
    count. drug is blank on every row.
  * the concentration unit. The S1b group headings are the bare numbers 12.5,
    25, 50 and 100. The numbers are recorded, conc_unit stays blank.
  * the organism. `UPEC`, `MG1655`, `UPEC_YFP` and `EHEC` are recorded as
    strain, as the sheets write them; no species is named anywhere.
  * plated volume, dilution, limit of detection -- absent from every sheet of
    every workbook (searched cell by cell).

INTERPRETIVE CALLS, each flagged on the rows it touches:

  1. FigS1 gives no unit for its values. They are recorded in cfu_per_ml on the
     reading that a viable-count kill curve running 1e7 to 1e2 over 240 minutes
     is per mL. If that is wrong, every cfu_per_ml here is wrong by one constant
     factor. So that the caveat survives into a structured column and not only
     into free text, readout on those rows reads `CFU (the sheet states no unit
     for these values)` rather than a bare `CFU`.
  2. Block S1c has no time header. Its grid (0, 30, 60, 120, 240) is the grid
     S1a states as minutes two blocks above it in the same sheet, and it is read
     as minutes on that basis.
  3. FigS4a heads its time column just `Time`. Sheet `c` of the same workbook
     heads its own `Time (min)`, and FigS4a's grid (0, 30, 60, 90, 120, 180,
     240, 360) is a minutes grid on that reading. Flagged on all 48 rows.

ONE DUPLICATION, verified rather than assumed: the four columns of the S1b `50`
group are cell-for-cell the same numbers as UPEC replicates 8 to 11 of S1a, at
the five timepoints the two blocks share; S1b carries two extra timepoints (180
and 360 min) that S1a does not. The reader checks this at run time and, when it
holds, says so in the notes of the 28 rows concerned so that nobody averages the
same four cultures in twice.
"""
from __future__ import annotations

import re
import warnings
from pathlib import Path

import openpyxl
import pandas as pd

from src.ingest import empty, finish

STUDY = "MIYAHARA2026"
SOURCE = "Fig1_S1.xlsx"
SHEET = "FigS1"
BLOCK = re.compile(r"^S1[a-z]$", re.I)
TIME_COL = 2

FLOOR_BASIS = ("no plated volume, dilution or limit of detection in any sheet "
               "of this deposit")
UNIT_NOTE = ("the sheet states no unit for these values; recorded as CFU/mL on "
             "the reading that the block is a viable-count kill curve")
COUNT_READOUT = "CFU (the sheet states no unit for these values)"


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def _txt(v) -> str:
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def read(d: Path) -> pd.DataFrame:
    path = d / SOURCE
    if not path.is_file():
        return empty()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        return empty()
    ws = wb[SHEET]

    g: dict[tuple[int, int], object] = {}
    for row in ws.iter_rows():
        for cell in row:
            g[(cell.row, cell.column)] = cell.value
    ncol = ws.max_column

    anchors = [r for r in range(1, ws.max_row + 1)
               if BLOCK.match(_txt(g.get((r, 1))))]
    if not anchors:
        return empty()

    # A time header anywhere in the sheet's time column, for the block that
    # gives none of its own.
    stated = [_txt(g.get((r, TIME_COL))) for r in anchors]
    minutes_from = next((s for s in stated if "min" in s.lower()), "")

    blocks = []
    for a in anchors:
        name = _txt(g.get((a, 1)))
        header = _txt(g.get((a, TIME_COL)))
        rows = []
        rr = a + 1
        while rr <= ws.max_row:
            t = _num(g.get((rr, TIME_COL)))
            if t is None:
                break
            rows.append((rr, t))
            rr += 1
        heads = [c for c in range(TIME_COL + 1, ncol + 1) if _txt(g.get((a, c)))]
        groups = []
        for i, h in enumerate(heads):
            end = heads[i + 1] - 1 if i + 1 < len(heads) else ncol
            cols = [c for c in range(h, end + 1)
                    if any(_num(g.get((r, c))) is not None for r, _ in rows)]
            if cols:
                groups.append((_txt(g.get((a, h))), cols))
        blocks.append({"name": name, "header": header, "rows": rows,
                       "groups": groups})

    # Is the S1b "50" group literally the same numbers as four S1a columns?
    def series(rows, col):
        """time -> value, so two blocks on different time grids still compare."""
        return {t: _num(g.get((r, col))) for r, t in rows
                if _num(g.get((r, col))) is not None}

    repeats: dict[tuple[str, int], str] = {}
    for i, b in enumerate(blocks):
        for gname, cols in b["groups"]:
            for col in cols:
                mine = series(b["rows"], col)
                if not mine:
                    continue
                for j, other in enumerate(blocks):
                    if j >= i:
                        continue
                    for oname, ocols in other["groups"]:
                        for ocol in ocols:
                            theirs = series(other["rows"], ocol)
                            shared = set(mine) & set(theirs)
                            if len(shared) >= 4 and all(
                                    mine[t] == theirs[t] for t in shared):
                                repeats[(b["name"], col)] = (
                                    f"identical values to replicate "
                                    f"{ocols.index(ocol) + 1} of "
                                    f"{other['name']} {oname!r} -- the same "
                                    f"culture reported twice, do not count both")

    out = []
    for b in blocks:
        header = b["header"]
        if "min" in header.lower():
            t_note = "source time in minutes"
            scale = 1 / 60.0
        elif minutes_from:
            t_note = (f"this block gives no time header; read as minutes from "
                      f"{minutes_from!r} on the same sheet")
            scale = 1 / 60.0
        else:
            t_note, scale = "no time unit stated on this sheet", None
        for gname, cols in b["groups"]:
            conc = _num(gname)
            strain = "" if conc is not None else gname
            for i, col in enumerate(cols, 1):
                for r, t_raw in b["rows"]:
                    v = _num(g.get((r, col)))
                    if v is None:
                        continue
                    note = "; ".join(x for x in [
                        t_note, UNIT_NOTE,
                        f"block {b['name']}, group heading {gname!r}",
                        ("group heading is a bare number; the sheet gives no "
                         "unit for it" if conc is not None else ""),
                        ("the block does not say which strain it used"
                         if conc is not None else ""),
                        repeats.get((b["name"], col), ""),
                    ] if x)
                    out.append({
                        "source_file": SOURCE, "sheet": SHEET,
                        "organism": "", "strain": strain, "drug": "",
                        "concentration": conc, "conc_unit": "",
                        "arm": f"{b['name']} {gname}",
                        "replicate": f"{b['name']}-{gname}-rep{i}",
                        "tech_replicate": "",
                        "time_h": t_raw * scale if scale else None,
                        "cfu_per_ml": v, "readout": COUNT_READOUT,
                        "floor_basis": FLOOR_BASIS, "notes": note,
                    })
    out.extend(_field_counts(d))
    out.extend(_survival_fraction(d))
    if not out:
        return empty()
    return finish(pd.DataFrame(out), STUDY)


# ---------------------------------------------------------------------------
# The two panels that are not counts but do say what they are.

def _open(d: Path, name: str):
    path = d / name
    if not path.is_file():
        return None
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path, data_only=True)


def _field_counts(d: Path) -> list[dict]:
    """Fig1b: `Time (min)` / `Cell count`, cells in a microscope field."""
    wb = _open(d, SOURCE)
    if wb is None or "Fig1b" not in wb.sheetnames:
        return []
    ws = wb["Fig1b"]
    if _txt(ws.cell(1, 1).value) != "Time (min)" or        _txt(ws.cell(1, 2).value) != "Cell count":
        return []                       # the sheet is not what it was
    rows = []
    for r in range(2, ws.max_row + 1):
        t, v = _num(ws.cell(r, 1).value), _num(ws.cell(r, 2).value)
        if t is None or v is None:
            continue
        rows.append({
            "source_file": SOURCE, "sheet": "Fig1b", "organism": "",
            "strain": "", "drug": "", "concentration": None, "conc_unit": "",
            "arm": "Fig1b cell count", "replicate": "Fig1b",
            "tech_replicate": "", "time_h": t / 60.0, "cfu_per_ml": None,
            "readout": "cell count in a microscope field of view",
            "floor_basis": ("not a colony count: the sheet counts cells in a "
                            "field of view, so no plating floor applies and "
                            "the deposit states none anyway"),
            "notes": (f"value={v:g} cells in the field of view; the schema has "
                      f"no column for a non-count reading, so the value is "
                      f"kept here and cfu_per_ml is left blank -- a field "
                      f"count is NOT a density and nothing converts it to one; "
                      f"the sheet heads the columns 'Time (min)' and 'Cell "
                      f"count'; source time in minutes; the sheet names no "
                      f"strain, organism or drug"),
        })
    return rows


def _survival_fraction(d: Path) -> list[dict]:
    """FigS4 sheet a: `Survival fraction`, six EHEC replicates, t=0 = 1."""
    wb = _open(d, "FigS4.xlsx")
    if wb is None or "a" not in wb.sheetnames:
        return []
    ws = wb["a"]
    if _txt(ws.cell(1, 2).value) != "Survival fraction":
        return []
    group = _txt(ws.cell(2, 2).value)               # 'EHEC'
    heads = {c: _txt(ws.cell(3, c).value)
             for c in range(2, ws.max_column + 1) if _txt(ws.cell(3, c).value)}
    if not heads:
        return []
    t_note = ("this sheet heads its time column just 'Time'; sheet 'c' of the "
              "same workbook heads its own 'Time (min)' and this grid (0, 30, "
              "60, 90, 120, 180, 240, 360) is read as minutes on that basis")
    rows = []
    for r in range(4, ws.max_row + 1):
        t = _num(ws.cell(r, 1).value)
        if t is None:
            continue
        for c, label in heads.items():
            v = _num(ws.cell(r, c).value)
            if v is None:
                continue
            rows.append({
                "source_file": "FigS4.xlsx", "sheet": "a", "organism": "",
                "strain": group, "drug": "", "concentration": None,
                "conc_unit": "", "arm": f"FigS4a {group}",
                "replicate": f"FigS4a-{label}", "tech_replicate": "",
                "time_h": t / 60.0, "cfu_per_ml": None,
                "readout": "CFU (surviving fraction only)",
                "floor_basis": ("no plated volume, dilution or limit of "
                                "detection in any sheet of this deposit; the "
                                "panel is a normalised fraction, so there is "
                                "no count to compare a floor with"),
                "notes": (f"value={v!r} surviving fraction; reported only as a "
                          f"fraction of its own t=0, which is written as 1, and "
                          f"the deposit gives no count at t=0, so this reading "
                          f"cannot be turned into CFU/mL; {t_note}; the sheet "
                          f"heads the block 'Survival fraction' over '{group}' "
                          f"and the column '{label}'; it names no drug and no "
                          f"organism"),
            })
    return rows
