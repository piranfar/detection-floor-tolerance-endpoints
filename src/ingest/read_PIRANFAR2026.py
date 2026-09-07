"""
Our own boundary experiment, read into the corpus schema.

E. coli ATCC 25922 against ciprofloxacin at 10x MIC, seeded at three densities
and plated at two volumes, measured at Iran University of Medical Sciences.

WHY IT BELONGS IN THE CORPUS RATHER THAN ONLY IN ITS OWN ANALYSIS. Almost none
of the 42 public deposits states the floor its own assay was working against --
that absence is what the corpus was built to count. This one does, by
construction: the plated volume is recorded per reading, so the floor is not
inferred, it is arithmetic. Putting it in the same table as the others makes it
the reference case, the one dataset where every quantity the analysis wants is
present, and it shows what the other rows are missing by contrast rather than by
assertion.

The floor here is one colony in the volume actually plated, carried back through
the dilution. Two technical plates of the same sample are pooled, so their
volumes add and the floor drops accordingly -- pooling buys measurable depth,
and getting that wrong is what made this project's own first analysis briefly
report that its boundary had been refuted.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.ingest import empty

ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "experiment" / "BOUNDARY_TEST_BLIND.xlsx"


def read(d: Path) -> pd.DataFrame:
    book = BOOK if BOOK.exists() else next(iter(d.glob("*.xlsx")), None)
    if book is None:
        return empty()
    ws = load_workbook(book, data_only=False)["Counts"]

    rows = []
    for r in range(5, ws.max_row + 1):
        v = [ws.cell(row=r, column=c).value for c in range(1, 8)]
        if not v[0] or v[6] is None:
            continue
        arm, conc, flask, t, vol, dil, col = v
        rows.append({
            "source_file": book.name,
            "sheet": "Counts",
            "organism": "Escherichia coli",
            "strain": "ATCC 25922",
            "drug": "ciprofloxacin",
            "concentration": float(conc) if conc is not None else None,
            "conc_unit": "xMIC",
            "arm": str(arm).strip(),
            "replicate": f"flask {str(flask).strip()}",
            "tech_replicate": "",
            "time_h": float(t),
            "colonies": float(col),
            "dilution": float(dil),
            "plated_volume_ul": float(vol),
            "readout": "CFU",
            "notes": "prospective test of the reachability and identifiability "
                     "boundaries; predictions recorded before counting",
        })
    df = pd.DataFrame(rows)
    if df.empty:
        return empty()

    # Pool the technical plates: volumes add, so the floor falls. Doing this at
    # ingest rather than downstream keeps one row per SAMPLE-plating, which is
    # the unit the rest of the corpus is in.
    keys = ["source_file", "sheet", "organism", "strain", "drug", "concentration",
            "conc_unit", "arm", "replicate", "time_h", "plated_volume_ul",
            "dilution", "readout", "notes"]
    g = df.groupby(keys, as_index=False, dropna=False).agg(
        colonies=("colonies", "sum"), n_plates=("colonies", "size"))
    g["plated_volume_ul"] = g.plated_volume_ul * g.n_plates
    g["tech_replicate"] = g.n_plates.map(lambda n: f"{n} plates pooled")
    g["floor_cfu_per_ml"] = 1000.0 * g.dilution / g.plated_volume_ul
    g["floor_basis"] = ("one colony in the pooled plated volume, carried back "
                        "through the dilution; the volume is recorded per reading")
    return g.drop(columns=["n_plates"])
