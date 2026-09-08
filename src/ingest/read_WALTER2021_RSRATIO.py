"""WALTER2021_RSRATIO -- Nature Communications Source Data (41467_2021_22833_MOESM4).

One workbook, ten sheets, of which seven carry viable counts: one in vitro
time-kill series (5 drugs plus a day-0 control, up to 9 sampling times), one in
vitro resistant-mutant series, and five BALB/c mouse experiments reporting CFU
per LUNG.  Three sheets carry no counts at all and are not read: 'Fig 1b & Sup
Fig 1a' (RS ratio and growth rate in an oxygen-depletion model), 'Fig 6a' and
'Fig 6b' (rRNA synthesis ratio and rifampin AUC in patient sputum).

WHAT THE DEPOSIT SAYS, AND WHAT IT DOES NOT

  * ORGANISM AND STRAIN ARE NEVER NAMED.  The full shared-string table of the
    workbook contains no species, no strain and no "H37Rv"; the closest it comes
    is "BALB/c chronic infection model" and "drug-susceptible TB" in a patient
    sheet that has no counts.  Both columns are therefore blank on every row.

  * CONCENTRATION IS NEVER GIVEN.  The in vitro arms are labelled only RIF, INH,
    STR, EMB, BDQ; the mouse arms carry a bare number ("RIF 10", "STR 200",
    "PZA 150 + RIF 10") whose unit the workbook never states.  A mouse dose is
    not a concentration in any case, so `concentration` and `conc_unit` are
    blank throughout and the label survives verbatim in `arm`.  The drug
    abbreviations and the regimen codes (2HRZE/3HR, PaMZ, BPaL, BPaMZ, HRZE) are
    expanded nowhere in the deposit and are passed through unexpanded.

  * IN VIVO COUNTS ARE PER LUNG, NOT PER mL.  Sheets 'Fig 1c', 'Fig 4a-c',
    'Fig 5a-d', 'Sup Fig 8b' and 'Sup Fig 11' head their count column "CFU per
    lung".  The value goes in cfu_per_ml because the schema has no per-organ
    column, and `readout` on every one of those rows says the denominator is a
    whole lung.  No volume is invented to convert it.

FLOORS.  Four of the seven sheets state one and three do not, and the split is
the point of reading this deposit at all.

  STATED, because the sheet itself writes the limit:
    - 'Fig 4a-c & Sup Fig 7-8', RESISTANT column: below-limit readings are
      written "< 2.18" and the smallest detected value is written
      2.176091259 = log10(150).  One colony is 150 CFU per lung, so the floor
      of that column is 150 CFU per lung.
    - 'Fig 5a-d & Sup Fig 9-10': below-limit readings are written "<3.26" and
      the smallest reported count is exactly 3.26 CFU per lung.
    - 'Sup Fig 8b': one reading is written "<1.6" (log10), floor 10**1.6.
    - 'Sup Fig 11': below-limit readings are written "<0.88" (log10) and 0.88
      is also written as a detected value, floor 10**0.88.
  NOT STATED, and left blank:
    - 'Fig 3a-e & Sup Fig 5', the in vitro time-kill sheet -- the one series a
      modeller would most want a floor for.  Nothing in the workbook gives a
      limit, a plated volume or a dilution.  The article states one in the
      legend to Supplementary Fig. 5, but that Supplementary PDF is NOT part of
      the deposited Source Data and is not in this deposit; the number is
      therefore not transcribed here, and a later pass that adds the PDF to the
      deposit should parse it rather than trust this note.
    - 'Sup Fig 5g-j' (resistant CFU per mL) and 'Fig 1c' (CFU per lung): no
      marker, no volume, no limit.
  Every stated floor is scoped to the sheet and column that states it.  The
  150 CFU per lung of the resistant column is NOT carried across to the total
  CFU column of the same sheet: Fig 1c reports untreated lungs at 1.85 log10 =
  71 CFU, below 150, so the two platings plainly do not share a limit.

BELOW-LIMIT READINGS carry no number in this workbook -- "<3.26" says only that
the count was under 3.26.  Those rows keep cfu_per_ml BLANK, are marked
censored = yes, and carry the floor; a censored-data model has everything it
needs and nothing has been made up.  The one place the deposit writes a bare 0
is the resistant column of 'Sup Fig 5g-j', which is a log10 column: a 0 there is
either 10**0 = 1 CFU/mL or "none detected", the sheet does not say which, and
those 29 readings are returned with a blank count and the ambiguity in notes.

TIME.  Each sheet gets its own conversion and each row says which was used.
  * 'Fig 3a-e' and 'Sup Fig 5g-j': "Treatment days" x 24 (day 0.25 = 6 h).
  * 'Fig 4a-c': treated rows are 28 days of treatment = 672 h.  This sheet's two
    time columns swap meaning between its baseline and its treated rows -- the
    baseline arms read (1, 0) and (11, 0), matching their own labels "Day 1 post
    infection" and "Pre-Rx", while the treated arms read (28, 39) with
    39 = 11 + 28.  'Sup Fig 11' prints the same pair of columns under the
    opposite headers, which is how the pairing was settled.
  * 'Fig 5a-d': "Day of sacrifice" x 24, counted from treatment start (it equals
    "Days of treatment" on treatment and "Days of treatment" + 84 at relapse).
  * 'Sup Fig 8b': its caption says 28-day treatment and every row shares
    "Day of sacrifice" 53, so all rows are 672 h.
  * 'Sup Fig 11': from the arm label, "(4 weeks)" = 672 h, "(8 weeks)" = 1344 h.
  * 'Fig 1c': untreated animals only.  time_h there is hours after AEROSOL
    INFECTION, not after drug exposure, and every row says so.
  Four kinds of row cannot be given a time and get a blank one rather than a
  guessed one: 'Fig 4a-c' "Day 1 post infection" (10 days before treatment
  started, so not time zero of treatment); the two 'Fig 4a-c' mice whose comment
  says "mouse euthanized early, day 18" / "day 17" while their Days-of-treatment
  column still reads 28, a contradiction the sheet does not resolve; and
  'Fig 5a-d' UNTX animals, whose day columns read "NA".

REPEATED READINGS.  Two are flagged rather than dropped: the day-0 Control of
'Fig 3a-e' Set 4 appears twice with one CFU value and two RS ratios, and the
three untreated mice of 'Sup Fig 8b' are the same animals, with the same
numbers, as the day-53 untreated mice of 'Fig 1c'.

NOT A READING, and skipped: 35 "not plated" cells in the resistant column of
'Fig 4a-c' and 4 "Not tested" cells in 'Sup Fig 5g-j'.  The 2 "Too numerous to
count" cells of 'Sup Fig 5g-j' and the 2 "nd" cells of 'Sup Fig 11' ARE kept,
with a blank count and the verbatim string in notes, because a reading was
attempted in each.
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.ingest import empty

BOOK = "_unpacked/PMC8131613_supplementary/41467_2021_22833_MOESM4_ESM.xlsx"

# ---------------------------------------------------------------- readouts
R_ML = "CFU per mL (the source's column is 'CFU per mL (log10)')"
R_ML_RES = ("resistant CFU per mL, colonies on drug-containing plates (the "
            "source's column is 'Resistant CFU per mL (log10)')")
R_LUNG = ("CFU per mouse lung (the source's column is 'CFU per lung'; the "
          "denominator is a whole lung, NOT a mL)")
R_LUNG_RES = ("resistant CFU per mouse lung, colonies on drug-containing plates "
              "(the source's column is 'Resistant CFU per lung (log10)'; the "
              "denominator is a whole lung, NOT a mL)")

# ---------------------------------------------------------------- floors
F_FIG4_RES = ('stated in the sheet: its "Resistant CFU per lung (log10)" column '
              'writes below-limit readings as "< 2.18" and writes its smallest '
              'detected value as 2.176091259, which is log10(150); the detected '
              'values are log10 of exact multiples of 150 (300, 600, 750, 2250), '
              'so one colony is 150 CFU per lung and 150 is this column\'s '
              'reporting floor')
F_FIG5 = ('stated in the sheet: its "CFU per lung" column writes below-limit '
          'readings as "<3.26" and its smallest reported count is exactly 3.26 '
          'CFU per lung, so one colony is 3.26 CFU per lung and 3.26 is this '
          'column\'s reporting floor')
F_S8B = ('stated in the sheet: its "CFU per lung (log10)" column writes a '
         'below-limit reading as "<1.6", so the floor is 10**1.6 = 39.8 CFU '
         'per lung')
F_S11 = ('stated in the sheet: its "CFU per lung (log10)" column writes '
         'below-limit readings as "<0.88" and writes 0.88 itself as a detected '
         'value, so the floor is 10**0.88 = 7.59 CFU per lung')

N_INVITRO = ('the workbook gives no detection limit, no plated volume and no '
             'dilution for the in vitro counts, and writes no below-limit '
             'marker in this column; the article states one in a supplementary '
             'figure legend, but that PDF is not part of the deposited Source '
             'Data and is not in this deposit, so no number is filled here')
N_FIG1C = ('this sheet gives no detection limit, plated volume or dilution for '
           'its "CFU per lung (log10)" column and writes no below-limit marker')
N_FIG4_TOT = ('this sheet\'s "CFU per lung (log10)" column carries no '
              'below-limit marker, no plated volume and no dilution; the floor '
              'its RESISTANT column states (150 CFU per lung) belongs to a '
              'different plating and is not carried over -- sheet "Fig 1c & Sup '
              'Fig 1b" reports untreated lungs at 1.85 log10 = 71 CFU, below '
              'that value')

# ---------------------------------------------------------------- notes
NO_ORG = ("the workbook names no organism and no strain anywhere, so both "
          "columns are blank")
NO_CONC = ("the workbook states no drug concentration and no dose unit; the "
           "arm label is the source's own and the drug abbreviations are "
           "expanded nowhere in the deposit")
DOSE_NOTE = ("the number in the arm label is the dose the sheet writes; the "
             "workbook never gives its unit, so it is not put in the "
             "concentration column")
BLANK_NOTE = ('the sheet gives no number for this reading, only that it is '
              'below the limit, so cfu_per_ml is left blank and the floor '
              'carries the information')

T_FIG3 = "time_h is the sheet's 'Treatment days' column x 24 (day 0.25 = 6 h)"
T_FIG1C = ("time_h is hours after AEROSOL INFECTION, not after drug exposure: "
           "every mouse in this sheet is untreated and the sheet's only time "
           "column is 'Day of sacrifice', counted from infection")
T_FIG4 = ("time_h is hours of treatment: this sheet's treated rows read 28 in "
          "'Days of treatment' and 39 in 'Day of sacrifice', and 39 - 28 = 11 "
          "is the day treatment began; its two baseline arms read (1, 0) and "
          "(11, 0), which match their own labels 'Day 1 post infection' and "
          "'Pre-Rx' only if the two columns swap meaning on those rows, and "
          "sheet 'Sup Fig 11' prints the same pair under the opposite headers")
T_FIG5 = ("time_h is this sheet's 'Day of sacrifice' x 24, counted from "
          "treatment start: it equals 'Days of treatment' for on-treatment rows "
          "and 'Days of treatment' + 84 for relapse rows")
T_S8B = ("time_h is 28 days x 24: this sheet's caption says '28-day treatment' "
         "and every row shares 'Day of sacrifice' 53, so treated and untreated "
         "animals were taken on the same day")
T_S11 = ("time_h is from the arm label ('4 weeks' = 28 d, '8 weeks' = 56 d), "
         "which matches the column this sheet prints under 'Day of sacrifice' "
         "(28, 56); this sheet's two time headers are swapped with respect to "
         "their values -- the column headed 'Days of treatment' holds 39 and "
         "67, which are 11 + 28 and 11 + 56, days from infection")

DIL_FIG5 = ("the counts in this column are whole-lung totals worked back from "
            "plates read at more than one dilution -- some blocks are multiples "
            "of 3.26 and others of 7.5 -- so 3.26 CFU per lung is the floor of "
            "the least diluted plating, and the effective floor of a diluted "
            "reading is higher and is not stated")

LOG0 = ('the sheet writes 0 in a log10 column; it does not say whether that '
        'means 10**0 = 1 CFU/mL or "none detected", so no count is recorded '
        'and the cell is reported here instead of being interpreted')


def _n(v) -> str:
    """A header cell, newlines and doubled spaces collapsed."""
    return re.sub(r"\s+", " ", str(v)).strip()


def _header(ws) -> tuple[int, dict]:
    """(row index, {header text -> column index}) for a sheet of this workbook."""
    for r in range(1, 12):
        vals = {_n(ws.cell(row=r, column=c).value): c
                for c in range(1, ws.max_column + 1)
                if ws.cell(row=r, column=c).value is not None}
        if "Mouse" in vals or "Set" in vals:
            return r, vals
    raise ValueError("%s: no header row" % ws.title)


def _caption(ws, hdr: int) -> str:
    """The sheet's own title lines, above the header row."""
    out = []
    for r in range(1, hdr):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip():
                out.append(_n(v))
    return "; ".join(out)


def _body(ws, hdr: int, key: int):
    """Data rows: those with something in the Mouse/Set column."""
    for r in range(hdr + 1, ws.max_row + 1):
        if ws.cell(row=r, column=key).value is not None:
            yield r


def _num(v):
    return float(v) if isinstance(v, (int, float)) and np.isfinite(v) else None


def _drug(arm: str) -> str:
    """The arm label with its dose figures stripped. '' where there is no drug."""
    a = arm.strip()
    if a.lower() in {"control", "untreated", "untx", "pre-rx", "pre-treatment",
                     "day 1 post infection"}:
        return ""
    parts = []
    for p in a.split("+"):
        p = re.sub(r"\s*\(.*?\)\s*", "", p)          # drop "(4 weeks)"
        parts.append(re.sub(r"\s*\d+(\.\d+)?\s*$", "", p).strip())  # drop "100"
    return " + ".join(p for p in parts if p)


def _join(*bits) -> str:
    return "; ".join(b for b in bits if b)


# --------------------------------------------------------------- the sheets

def _fig1c(ws) -> list[dict]:
    """Untreated BALB/c chronic infection: CFU per lung, days 1-53 after infection."""
    hdr, col = _header(ws)
    cap = _caption(ws, hdr)
    out = []
    for r in _body(ws, hdr, col["Mouse"]):
        mouse = ws.cell(row=r, column=col["Mouse"]).value
        arm = _n(ws.cell(row=r, column=col["Treatment"]).value)
        day = _num(ws.cell(row=r, column=col["Day of sacrifice"]).value)
        v = _num(ws.cell(row=r, column=col["CFU per lung (log10)"]).value)
        if v is None:
            continue
        dup = ("this mouse and its numbers appear again on sheet 'Sup Fig 8b' "
               "as one of that sheet's untreated controls" if day == 53 else "")
        out.append({
            "sheet": ws.title, "arm": arm, "drug": _drug(arm),
            "replicate": "fig1c mouse %s" % mouse,
            "time_h": None if day is None else day * 24.0,
            "cfu_per_ml": 10.0 ** v,
            "floor_basis": N_FIG1C, "readout": R_LUNG,
            "notes": _join(cap, T_FIG1C, dup, NO_ORG, NO_CONC,
                           "sheet value %r log10 CFU per lung" % v),
        })
    return out


def _fig3(ws) -> list[dict]:
    """In vitro time-kill: 5 drugs plus a day-0 control, CFU per mL (log10)."""
    hdr, col = _header(ws)
    cap = _caption(ws, hdr)
    out, seen = [], set()
    for r in _body(ws, hdr, col["Set"]):
        st = ws.cell(row=r, column=col["Set"]).value
        arm = _n(ws.cell(row=r, column=col["Group"]).value)
        day = _num(ws.cell(row=r, column=col["Treatment days"]).value)
        v = _num(ws.cell(row=r, column=col["CFU per mL (log10)"]).value)
        if v is None or day is None:
            continue
        k = (st, arm, day, v)
        dup = ("this sheet writes the same reading twice, on two rows that "
               "differ only in their RS ratio; both are kept and this is the "
               "repeat" if k in seen else "")
        seen.add(k)
        ctrl = ("'Control' appears in this sheet only at treatment day 0: it is "
                "the pre-treatment reading, not an untreated arm followed over "
                "time" if arm.lower() == "control" else "")
        out.append({
            "sheet": ws.title, "arm": arm, "drug": _drug(arm),
            "replicate": "fig3 set %s" % st,
            "time_h": day * 24.0, "cfu_per_ml": 10.0 ** v,
            "floor_basis": N_INVITRO, "readout": R_ML,
            "notes": _join(cap, T_FIG3, ctrl, dup, NO_ORG, NO_CONC,
                           "sheet value %r log10 CFU per mL" % v),
        })
    return out


def _sup5gj(ws) -> list[dict]:
    """In vitro resistant mutants: 'Resistant CFU per mL (log10)', days 0-8."""
    hdr, col = _header(ws)
    cap = _caption(ws, hdr)
    out = []
    for r in _body(ws, hdr, col["Set"]):
        st = ws.cell(row=r, column=col["Set"]).value
        arm = _n(ws.cell(row=r, column=col["Group"]).value)
        day = _num(ws.cell(row=r, column=col["Treatment days"]).value)
        raw = ws.cell(row=r, column=col["Resistant CFU per mL (log10)"]).value
        cfu, extra = None, ""
        if isinstance(raw, str):
            if _n(raw).lower() == "not tested":
                continue                      # no reading was taken
            extra = 'the sheet writes this reading as "%s"' % _n(raw)
        else:
            v = _num(raw)
            if v is None:
                continue
            if v == 0:
                extra = LOG0
            else:
                cfu = 10.0 ** v
        out.append({
            "sheet": ws.title, "arm": arm, "drug": _drug(arm),
            "replicate": "sup5gj set %s" % st,
            "time_h": None if day is None else day * 24.0,
            "cfu_per_ml": cfu,
            "floor_basis": N_INVITRO, "readout": R_ML_RES,
            "notes": _join(cap, T_FIG3, extra, NO_ORG, NO_CONC,
                           "sheet value %r" % raw),
        })
    return out


def _fig4(ws) -> list[dict]:
    """BALB/c, 28 days of monotherapy: total and resistant CFU per lung."""
    hdr, col = _header(ws)
    cap = _caption(ws, hdr)
    ctot = col["CFU per lung (log10)"]
    cres = col["Resistant CFU per lung (log10)"]
    out = []
    for r in _body(ws, hdr, col["Mouse"]):
        mouse = ws.cell(row=r, column=col["Mouse"]).value
        arm = _n(ws.cell(row=r, column=col["Treatment"]).value)
        d_treat = _num(ws.cell(row=r, column=col["Days of treatment"]).value)
        d_sac = _num(ws.cell(row=r, column=col["Day of sacrifice"]).value)
        comment = ws.cell(row=r, column=col["CFU comment"]).value
        comment = _n(comment) if isinstance(comment, str) else ""

        when, why = None, ""
        if arm == "Pre-Rx":
            when = 0.0
            why = ("this is the pre-treatment reading, day 0 of treatment; the "
                   "sheet's own columns read 11 and 0, the 11 being the day "
                   "after infection on which treatment started")
        elif arm == "Day 1 post infection":
            why = ("no time: this reading is 10 days BEFORE treatment started "
                   "(day 1 after infection, treatment began on day 11), so it "
                   "is not time zero of treatment and no negative time is "
                   "invented for it")
        elif "euthanized early" in comment:
            why = ("no time: the sheet's 'Days of treatment' column reads %s for "
                   "this mouse while its own comment says it was euthanized "
                   "early, and the sheet does not say whether that day is "
                   "counted from infection or from treatment start, so the two "
                   "disagree and neither is used" % d_treat)
        elif d_treat is not None:
            when = d_treat * 24.0
            if d_sac is not None and abs(d_sac - d_treat - 11) > 1e-9:
                why = ("this row's 'Day of sacrifice' is %s, which is not "
                       "'Days of treatment' + 11 as it is on every other "
                       "treated row of the sheet" % d_sac)

        base = {"sheet": ws.title, "arm": arm, "drug": _drug(arm),
                "replicate": "fig4 mouse %s" % mouse, "time_h": when}
        common = _join(cap, T_FIG4, why, comment and
                       'the sheet\'s CFU comment column says "%s"' % comment,
                       NO_ORG, NO_CONC, DOSE_NOTE)

        v = _num(ws.cell(row=r, column=ctot).value)
        if v is not None:
            out.append({**base, "cfu_per_ml": 10.0 ** v,
                        "floor_basis": N_FIG4_TOT, "readout": R_LUNG,
                        "notes": _join(common,
                                       "sheet value %r log10 CFU per lung" % v)})

        raw = ws.cell(row=r, column=cres).value
        cfu, cens, extra = None, "", ""
        if isinstance(raw, str):
            if _n(raw).lower() == "not plated":
                continue                      # the assay was not done
            cfu, cens = None, "yes"
            extra = _join('the sheet writes this reading as "%s"' % _n(raw),
                          BLANK_NOTE)
        else:
            v = _num(raw)
            if v is None:
                continue
            cfu = 10.0 ** v
            if abs(v - 2.176091259) < 1e-6:
                extra = ("this is the sheet's censoring value itself, one "
                         "colony = 150 CFU per lung: a detection at the floor, "
                         "not a below-limit reading")
        out.append({**base, "cfu_per_ml": cfu, "censored": cens,
                    "floor_cfu_per_ml": 10.0 ** 2.176091259,
                    "floor_basis": F_FIG4_RES, "readout": R_LUNG_RES,
                    "notes": _join(common, extra, "sheet value %r" % raw)})
    return out


def _fig5(ws) -> list[dict]:
    """The relapsing mouse model: CFU per lung, linear, 16 sacrifice days."""
    hdr, col = _header(ws)
    cap = _caption(ws, hdr)
    out = []
    for r in _body(ws, hdr, col["Mouse"]):
        mouse = ws.cell(row=r, column=col["Mouse"]).value
        arm = _n(ws.cell(row=r, column=col["Treatment"]).value)
        stage = _n(ws.cell(row=r, column=col["Stage at which sacrificed"]).value)
        d_treat = ws.cell(row=r, column=col["Days of treatment"]).value
        d_sac = _num(ws.cell(row=r, column=col["Day of sacrifice"]).value)
        raw = ws.cell(row=r, column=col["CFU per lung"]).value
        if raw is None:
            continue

        when, why = None, ""
        if d_sac is not None:
            when = d_sac * 24.0
        elif stage == "Pre-Treatment":
            when = 0.0
            why = ("the sheet's day columns read 'NA' here; this arm is labelled "
                   "'Pre-Rx' at the 'Pre-Treatment' stage, so it is day 0 of "
                   "treatment")
        else:
            why = ("no time: the sheet writes 'NA' in both day columns for these "
                   "untreated animals and says nowhere when they were taken")

        cfu, cens, extra = None, "", ""
        if isinstance(raw, str):
            cfu, cens = None, "yes"
            extra = _join('the sheet writes this reading as "%s"' % _n(raw),
                          BLANK_NOTE)
        else:
            cfu = _num(raw)
            if cfu is None:
                continue
        out.append({
            "sheet": ws.title, "arm": arm, "drug": _drug(arm),
            "replicate": "fig5 mouse %s" % mouse, "time_h": when,
            "cfu_per_ml": cfu, "censored": cens,
            "floor_cfu_per_ml": 3.26, "floor_basis": F_FIG5,
            "readout": R_LUNG,
            "notes": _join(cap, T_FIG5, why,
                           "the sheet's stage for this animal is '%s'" % stage,
                           "the sheet's 'Days of treatment' cell reads %r"
                           % d_treat,
                           DIL_FIG5, NO_ORG, NO_CONC, DOSE_NOTE,
                           "sheet value %r CFU per lung" % raw),
        })
    return out


def _sup8b(ws) -> list[dict]:
    """Low-dose aerosol BALB/c, 28-day treatment: CFU per lung (log10)."""
    hdr, col = _header(ws)
    cap = _caption(ws, hdr)
    out = []
    for r in _body(ws, hdr, col["Mouse"]):
        mouse = ws.cell(row=r, column=col["Mouse"]).value
        arm = _n(ws.cell(row=r, column=col["Treatment"]).value)
        d_sac = _num(ws.cell(row=r, column=col["Day of sacrifice"]).value)
        raw = ws.cell(row=r, column=col["CFU per lung (log10)"]).value
        if raw is None:
            continue
        cfu, cens, extra = None, "", ""
        if isinstance(raw, str):
            cfu, cens = None, "yes"
            extra = _join('the sheet writes this reading as "%s"' % _n(raw),
                          BLANK_NOTE)
        else:
            v = _num(raw)
            if v is None:
                continue
            cfu = 10.0 ** v
        untx = ("this animal was never treated; the 672 h is the day the treated "
                "animals of this sheet reached, which it shares; the same mouse, "
                "with the same numbers, is also the day-53 untreated animal of "
                "sheet 'Fig 1c & Sup Fig 1b'" if _drug(arm) == "" else "")
        out.append({
            "sheet": ws.title, "arm": arm, "drug": _drug(arm),
            "replicate": "sup8b mouse %s" % mouse,
            "time_h": 28 * 24.0, "cfu_per_ml": cfu, "censored": cens,
            "floor_cfu_per_ml": 10.0 ** 1.6, "floor_basis": F_S8B,
            "readout": R_LUNG,
            "notes": _join(cap, T_S8B, untx, extra,
                           "the sheet's 'Day of sacrifice' cell reads %r" % d_sac,
                           NO_ORG, NO_CONC, DOSE_NOTE,
                           "sheet value %r" % raw),
        })
    return out


def _sup11(ws) -> list[dict]:
    """BALB/c, 4 or 8 weeks of HRZE or BPaL: CFU per lung (log10)."""
    hdr, col = _header(ws)
    cap = _caption(ws, hdr)
    # This sheet's two time headers are swapped with respect to their values;
    # the arm label carries the treatment length, so use it and check it.
    out = []
    for r in _body(ws, hdr, col["Mouse"]):
        mouse = ws.cell(row=r, column=col["Mouse"]).value
        arm = _n(ws.cell(row=r, column=col["Treatment"]).value)
        printed = _num(ws.cell(row=r, column=col["Day of sacrifice"]).value)
        other = _num(ws.cell(row=r, column=col["Days of treatment"]).value)
        raw = ws.cell(row=r, column=col["CFU per lung (log10)"]).value
        if raw is None:
            continue
        m = re.search(r"\((\d+)\s*weeks?\)", arm)
        when = float(m.group(1)) * 7 * 24.0 if m else None
        why = ""
        if when is not None and printed is not None and \
                abs(printed * 24.0 - when) > 1e-9:
            why = ("the arm label and the sheet's own day columns disagree for "
                   "this row (label %s, columns %s and %s)" % (arm, printed, other))
        cfu, cens, extra = None, "", ""
        if isinstance(raw, str):
            s = _n(raw)
            if s.startswith("<"):
                cfu, cens = None, "yes"
                extra = _join('the sheet writes this reading as "%s"' % s,
                              BLANK_NOTE)
            else:
                extra = ('the sheet writes this reading as "%s" and does not say '
                         'what it stands for, so no count is recorded' % s)
        else:
            v = _num(raw)
            if v is None:
                continue
            cfu = 10.0 ** v
            if abs(v - 0.88) < 1e-9:
                extra = ("this is the sheet's censoring value itself: a "
                         "detection at the floor, not a below-limit reading")
        out.append({
            "sheet": ws.title, "arm": arm, "drug": _drug(arm),
            "replicate": "sup11 mouse %s" % mouse,
            "time_h": when, "cfu_per_ml": cfu, "censored": cens,
            "floor_cfu_per_ml": 10.0 ** 0.88, "floor_basis": F_S11,
            "readout": R_LUNG,
            "notes": _join(cap, T_S11, why, extra,
                           "this sheet's day cells read %r under 'Day of "
                           "sacrifice' and %r under 'Days of treatment'"
                           % (printed, other),
                           NO_ORG, NO_CONC, DOSE_NOTE,
                           "sheet value %r" % raw),
        })
    return out


SHEETS = {
    "Fig 1c & Sup Fig 1b": _fig1c,
    "Fig 3a-e & Sup Fig 5": _fig3,
    "Sup Fig 5g-j": _sup5gj,
    "Fig 4a-c & Sup Fig 7-8": _fig4,
    "Fig 5a-d & Sup Fig 9-10": _fig5,
    "Sup Fig 8b": _sup8b,
    "Sup Fig 11": _sup11,
}

# no bacterial counts in these; kept here so the omission is deliberate
NOT_READ = ("'Fig 1b & Sup Fig 1a' (RS ratio, ITS1/23S and growth rate in an "
            "oxygen-depletion model), 'Fig 6a' and 'Fig 6b' (patient sputum "
            "rRNA synthesis ratio and rifampin AUC) hold no bacterial counts")


def read(d: Path) -> pd.DataFrame:
    book = d / BOOK
    if not book.exists():
        return empty()
    wb = load_workbook(book, data_only=True, read_only=False)

    rows = []
    for name, fn in SHEETS.items():
        if name not in wb.sheetnames:
            continue
        rows.extend(fn(wb[name]))
    if not rows:
        return empty()

    df = pd.DataFrame(rows)
    df["source_file"] = BOOK
    df["organism"] = ""
    df["strain"] = ""
    df["notes"] = df["notes"] + ("; not read from this workbook: " + NOT_READ)
    return df
