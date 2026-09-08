"""
What a second method finds in lungs the plate calls sterile.

Run:  python -m src.experiments.exp43_below_the_floor_is_not_zero

WHY THIS EXISTS. The manuscript argues that a reading at the assay floor is not
a measurement of the surviving population, and that the region beneath the floor
is where relapse comes from. Both are inferences from arithmetic and from
crossing-and-returning series. Neither is a direct observation of organisms the
plate missed, because a plate cannot report what it cannot see -- that is the
whole problem, and it makes the claim hard to demonstrate rather than merely
argue.

Evangelopoulos and colleagues demonstrated it, on the way to a different point.
They enumerated the SAME mouse lungs three ways: colony counts, a molecular
bacterial load from 16S rRNA, and a most probable number. Under the deepest
regimens the colony count reads zero in every animal while the other two read
thousands per lung.

THE MPN IS THE ONE THAT MATTERS, and the distinction is worth stating carefully.
A molecular load counts nucleic acid, and nucleic acid survives its owner, so a
sceptic can attribute MBL signal to dead organisms. An MPN is a CULTURE: it
scores growth in replicate dilutions. Growth requires a viable, culturable
organism. So where the plate reads zero and the MPN reads thousands, the animal
carried thousands of organisms that were alive, culturable, and invisible to the
plate that declared the lung sterile.

That is this paper's claim, observed rather than derived, in a published
CC-BY-4.0 deposit, in the mouse lung model used to decide which tuberculosis
regimens enter clinical trials.

WHAT THIS DEPOSIT CANNOT DO, and is not asked to. It carries no time series and
no pre-treatment count, so no headroom is computable from it and it is absent
from the boundary sweep. It is used here for one thing only.

Writes:
  results/tables/exp43_plate_zero_vs_mpn.csv
  results/receipts/exp43_receipt.json
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BOOK = (ROOT / "data" / "corpus" / "open" / "EVANGELOPOULOS2022"
        / "34079879.xlsx")
TABLES = ROOT / "results" / "tables"
RECEIPTS = ROOT / "results" / "receipts"


def read_blocks(ws) -> dict:
    """Pull every (arm, assay) column out of a sheet laid out under 'Assay'.

    The workbook puts two independent blocks side by side on some sheets, each
    with its own 'Assay' header and its own arm labels, so the sheet is scanned
    for header cells rather than assumed to have one table.
    """
    data: dict[tuple[str, str], list[float]] = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(row=r, column=c).value).strip().lower() != "assay":
                continue
            arms, cc = [], c + 1
            while cc <= ws.max_column:
                v = ws.cell(row=r, column=cc).value
                if v is None or not str(v).strip():
                    break
                arms.append((str(v).strip(), cc))
                cc += 1
            rr = r + 1
            while rr <= ws.max_row:
                assay = ws.cell(row=rr, column=c).value
                if assay is None or not str(assay).strip():
                    break
                for arm, col in arms:
                    v = ws.cell(row=rr, column=col).value
                    if isinstance(v, (int, float)):
                        data.setdefault((arm, str(assay).strip()), []).append(float(v))
                rr += 1
    return data


def main() -> int:
    if not BOOK.exists():
        raise SystemExit(f"missing {BOOK}")
    TABLES.mkdir(parents=True, exist_ok=True)
    RECEIPTS.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(BOOK, data_only=True)

    rows = []
    for sheet in wb.sheetnames:
        d = read_blocks(wb[sheet])
        for arm in sorted({a for a, _ in d}):
            cfu = d.get((arm, "CFU"), [])
            mpn = d.get((arm, "MPN"), [])
            mbl = d.get((arm, "MBL"), [])
            if not cfu:
                continue
            rows.append({
                "sheet": sheet, "arm": arm,
                "n_animals": len(cfu),
                "n_cfu_zero": sum(1 for x in cfu if x == 0),
                "cfu_max": max(cfu),
                "n_mpn": len(mpn),
                "mpn_min": min(mpn) if mpn else None,
                "mpn_max": max(mpn) if mpn else None,
                "n_mbl": len(mbl),
                "mbl_min": min(mbl) if mbl else None,
                "mbl_max": max(mbl) if mbl else None,
            })
    t = pd.DataFrame(rows)
    t.to_csv(TABLES / "exp43_plate_zero_vs_mpn.csv", index=False)

    print(f"{len(t)} arms across {t.sheet.nunique()} sheets\n")
    print(f"   {'sheet':10}{'arm':24}{'mice':>5}{'CFU=0':>7}{'CFU max':>10}"
          f"{'MPN min':>9}{'MPN max':>9}")
    for r in t.itertuples():
        print(f"   {r.sheet:10}{r.arm[:24]:24}{r.n_animals:>5}{r.n_cfu_zero:>7}"
              f"{r.cfu_max:>10,.0f}"
              f"{r.mpn_min if r.mpn_min is not None else 0:>9,.0f}"
              f"{r.mpn_max if r.mpn_max is not None else 0:>9,.0f}")

    # The claim, isolated: arms where the plate found nothing in any animal AND
    # a culture-based MPN was run on the same lungs.
    sterile = t[(t.n_cfu_zero == t.n_animals) & (t.n_mpn > 0)]
    print(f"\n-- arms where the plate read zero in EVERY animal, and an MPN was "
          f"run on the same lungs --")
    if sterile.empty:
        print("   none")
    else:
        for r in sterile.itertuples():
            print(f"   {r.sheet} {r.arm}: {r.n_animals} of {r.n_animals} lungs "
                  f"sterile by plate count, MPN {r.mpn_min:,.0f} to "
                  f"{r.mpn_max:,.0f} per lung")
        print(f"\n   {int(sterile.n_animals.sum())} animals in total. The MPN is "
              f"a culture, not a nucleic acid,")
        print(f"   so the organisms it counts were alive and culturable. The "
              f"lowest MPN recorded")
        print(f"   in any of them is {sterile.mpn_min.min():,.0f} per lung.")

    (RECEIPTS / "exp43_receipt.json").write_text(json.dumps({
        "script": "src/experiments/exp43_below_the_floor_is_not_zero.py",
        "utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "deposit": "EVANGELOPOULOS2022, figshare 10.5522/04/19175153.v2, CC BY 4.0",
        "source_file": BOOK.name,
        "n_arms": int(len(t)),
        "arms_sterile_by_plate_with_mpn": sterile.arm.tolist(),
        "n_animals_sterile_by_plate_with_mpn": int(sterile.n_animals.sum())
        if not sterile.empty else 0,
        "lowest_mpn_in_a_plate_sterile_lung": float(sterile.mpn_min.min())
        if not sterile.empty else None,
        "highest_mpn_in_a_plate_sterile_lung": float(sterile.mpn_max.max())
        if not sterile.empty else None,
        "note": ("This deposit carries no time series and no pre-treatment "
                 "count, so no headroom is computable from it and it does not "
                 "enter the boundary sweep. It is used for one observation."),
    }, indent=2, default=str), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
