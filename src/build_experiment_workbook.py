"""
The data-capture workbook for the prospective test of the two boundaries.

Run:  python -m src.build_experiment_workbook

WHAT THE EXPERIMENT HAS TO DO. The manuscript's claim is arithmetic: a q-log
endpoint is observable only in a culture seeded above L*10^q, and above L/c1 a
floor-censored reading is compatible with only the lowest tolerance class. Read
off five published deposits, that is a reanalysis. Demonstrated forward, on a
culture whose starting density you set on purpose, it is a result.

The design here is built around one point that makes the test unusually cheap and
unusually clean: THE FLOOR CAN BE MOVED WITHOUT TOUCHING THE BIOLOGY. Plating
100 uL puts the smallest reportable positive count at 10 CFU/mL; plating 10 uL
puts it at 100. Same flask, same drug, same cells, same moment -- two floors, and
therefore two different deepest reportable endpoints. Any difference in the
tolerance call between those two platings cannot be biology, because there is only
one culture. That is the control an inoculum-only design does not have, and it is
why the workbook asks for every sample to be plated at two volumes.

The inoculum arm is the second half: seed the same strain low, middle and high so
that one level sits below L*10^4 and one above it, and the prediction is that the
low-seeded cultures cannot report a four-log reduction however hard they are
killed.

WHAT THE WORKBOOK ENFORCES. The predictions are computed on the Design sheet from
the numbers you enter there, BEFORE any count exists, and the Prediction sheet
locks them in words. The Counts sheet is where data goes afterwards. Keeping them
in that order in one file is the cheapest form of pre-specification available
without a registry, and it is what lets the analysis say the boundaries were
fixed in advance rather than fitted.

Writes: experiment/TIMEKILL_BOUNDARY_TEST.xlsx
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "experiment" / "TIMEKILL_BOUNDARY_TEST.xlsx"

INK = "1A1A1A"
MUTED = "6B6B6B"
HEAD_BG = "1F3B57"
ENTRY_BG = "FFF8E1"
CALC_BG = "EEF3F8"
NOTE_BG = "F5F5F2"
RULE = Side(style="thin", color="C8C8C0")
BOX = Border(left=RULE, right=RULE, top=RULE, bottom=RULE)

H1 = Font(bold=True, size=14, color=INK)
H2 = Font(bold=True, size=11, color=INK)
TH = Font(bold=True, size=10, color="FFFFFF")
BODY = Font(size=10, color=INK)
SMALL = Font(size=9, color=MUTED)
MONO = Font(size=10, name="Consolas", color=INK)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.border = TH, BOX
        cell.fill = PatternFill("solid", fgColor=HEAD_BG)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def note(ws, row, text, span=8, size=9):
    ws.cell(row=row, column=1, value=text).font = Font(size=size, color=MUTED)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    ws.cell(row=row, column=1).alignment = WRAP
    return row + 1


def readme(wb):
    ws = wb.create_sheet("Read me first")
    ws.column_dimensions["A"].width = 108
    r = 1
    ws.cell(row=r, column=1, value="Prospective test of the reachability and "
            "identifiability boundaries").font = H1
    r += 2
    for para in [
        "WHAT THIS TESTS. A q-log endpoint can only be observed in a culture that "
        "starts far enough above the assay floor to have q logs to fall. Write L "
        "for the smallest positive count the method can report and N0 for the "
        "starting density; the headroom is h = log10(N0/L), and a q-log endpoint "
        "is reportable only where h >= q. The second boundary is about labels: "
        "above L/c1, where c1 is the lowest tolerance-class threshold, a reading "
        "that lands on the floor is compatible with only the lowest class, "
        "whatever the drug did.",

        "READ THIS BEFORE SPENDING ANY MONEY. Half of what this workbook was built "
        "to test has already been measured by somebody else, and the data are "
        "public. The ERA4TB six-laboratory deposit plates every sample four ways -- "
        "100 uL, 10 uL twice, and 2.5 uL -- which is three different floors on one "
        "culture. In 20 of its 72 treated flasks, 27.8 per cent, those four "
        "platings disagree about whether that culture ever fell below the limit. "
        "One flask, one drug, one moment, four answers, decided by the pipette. "
        "That is the floor-moving control, already done, by six independent "
        "laboratories under one written protocol, at no cost. Do not buy it again.",

        "WHAT IS ACTUALLY WORTH BUYING is the other arm, and only that arm. The "
        "ERA4TB laboratories did not vary the INOCULUM on purpose: their starting "
        "densities differ, but the difference is confounded with which laboratory "
        "did the work, so it cannot separate the reachability boundary from every "
        "other thing that differs between laboratories. Seeding one strain, in one "
        "laboratory, with one drug, at three deliberate densities does separate it. "
        "That is this workbook, and it is half the experiment it started as.",

        "THE PREDICTION. Seed at three densities chosen so one sits below L*10^4 "
        "and one comfortably above. The low-seeded cultures cannot report a "
        "four-log reduction however completely they are killed, and their reported "
        "surviving fraction converges on L/N0 rather than on anything the drug did. "
        "The defaults on the Design sheet are chosen so the MIDDLE level is the "
        "5x10^5 inoculum the standard itself specifies -- which is legal at the "
        "100 uL plating and impossible at 10 uL. Keep plating both ways: it costs "
        "one extra plate and it makes each culture its own control.",

        "FILL IN THE DESIGN SHEET FIRST, AND DO NOT CHANGE IT AFTERWARDS. The "
        "predictions on the Prediction sheet are computed from it. Recording them "
        "before the first count is what allows the paper to say the boundaries were "
        "fixed in advance rather than fitted to the result. If you have to change "
        "the design after seeing data, save a new copy of this file and say so.",

        "YELLOW cells are for you to fill. BLUE cells compute themselves -- do not "
        "type in them.",
    ]:
        c = ws.cell(row=r, column=1, value=para)
        c.font, c.alignment = BODY, WRAP
        ws.row_dimensions[r].height = 15 * (1 + len(para) // 95)
        r += 2

    ws.cell(row=r, column=1, value="The order to work in").font = H2
    r += 1
    for i, step in enumerate([
        "Design sheet: strain, drug, concentrations, the three target inocula, "
        "the plated volumes, and the tolerance-class thresholds you will score "
        "against.",
        "Prediction sheet: read it, and check you are willing to be held to it. "
        "It follows from the Design sheet alone.",
        "Counts sheet: one row per plate. Enter colonies counted, the dilution and "
        "the volume plated; CFU/mL and the below-floor flag compute themselves.",
        "Day 0 check: confirm the realised N0 is close to the target. If it is "
        "not, the design still works -- the boundaries use the realised value -- "
        "but record what you actually got.",
        "When the counts are complete, hand the file back and the analysis runs "
        "against the predictions already written down.",
    ], 1):
        c = ws.cell(row=r, column=1, value=f"{i}.  {step}")
        c.font, c.alignment = BODY, WRAP
        ws.row_dimensions[r].height = 15 * (1 + len(step) // 95)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="What would refute the claim").font = H2
    r += 1
    for para in [
        "If low-seeded and high-seeded cultures return the SAME reported surviving "
        "fraction at the deep endpoint, the boundary is wrong.",
        "If the tolerance class assigned from the 100 uL plating agrees with the "
        "class from the 10 uL plating in every floor-censored sample, the "
        "identifiability boundary does not bite at these settings.",
        "Both are recorded as outcomes, not as failures. A boundary that can be "
        "refuted and is not is worth more than one that cannot be.",
    ]:
        c = ws.cell(row=r, column=1, value="   " + para)
        c.font, c.alignment = BODY, WRAP
        ws.row_dimensions[r].height = 15 * (1 + len(para) // 95)
        r += 1
    return ws


def design(wb):
    ws = wb.create_sheet("Design")
    ws.cell(row=1, column=1, value="Design — fill this in BEFORE any counting").font = H1
    r = 3
    header(ws, r, ["What", "Value", "Unit", "Note"], [34, 18, 16, 62])
    r += 1
    rows = [
        ("Strain", "Escherichia coli ATCC 25922", "", "the CLSI quality-control reference strain: the one the standard itself names"),
        ("Drug", "ciprofloxacin", "", "kills deeply and fast, so the floor is reached within hours"),
        ("Concentration tested", "10", "x MIC", "deep enough to reach the floor in every arm"),
        ("MIC of this strain", "", "mg/L", "measure it; do not take it from a table"),
        ("Sampling times", "0, 1, 2, 4, 6, 24", "hours", "hour 0 is essential: it is N0. E. coli, so hours not days"),
        ("Biological replicates per arm", "3", "flasks", "three is the minimum that gives a spread"),
        ("Technical plates per sample", "2", "plates", "per plated volume"),
        ("", "", "", ""),
        ("PLATED VOLUME A", "100", "uL", "the sensitive plating"),
        ("PLATED VOLUME B", "10", "uL", "the insensitive plating — same sample, higher floor"),
        ("", "", "", ""),
        ("Target inoculum LOW", "50000", "CFU/mL", "5e4: h = 3.7 at 100 uL and 2.7 at 10 uL, so a 4-log call is IMPOSSIBLE either way"),
        ("Target inoculum MID", "500000", "CFU/mL", "5e5, the CLSI standard inoculum: h = 4.7 at 100 uL but 3.7 at 10 uL. It STRADDLES"),
        ("Target inoculum HIGH", "5000000", "CFU/mL", "5e6: h = 5.7 and 4.7, so a 4-log call is legal either way"),
        ("", "", "", ""),
        ("Lowest class threshold c1", "0.001", "fraction", "e.g. 10^-3: below this the call is 'low tolerance'"),
        ("Middle class threshold c2", "0.01", "fraction", "e.g. 10^-2"),
        ("Deepest endpoint scored q", "4", "logs", "4 = a 99.99% reduction"),
    ]
    first_entry = r
    for what, val, unit, n in rows:
        ws.cell(row=r, column=1, value=what).font = H2 if what.isupper() else BODY
        c = ws.cell(row=r, column=2, value=val)
        c.font, c.border = BODY, BOX
        if what:
            c.fill = PatternFill("solid", fgColor=ENTRY_BG)
        ws.cell(row=r, column=3, value=unit).font = SMALL
        d = ws.cell(row=r, column=4, value=n)
        d.font, d.alignment = SMALL, WRAP
        r += 1

    volA, volB = first_entry + 8, first_entry + 9
    lo, mid, hi = first_entry + 11, first_entry + 12, first_entry + 13
    c1, q = first_entry + 15, first_entry + 17

    r += 1
    ws.cell(row=r, column=1, value="What those choices imply — computed, do not edit").font = H2
    r += 1
    header(ws, r, ["Quantity", "Plating A (100 uL)", "Plating B (10 uL)", "Meaning"],
           [34, 18, 18, 62])
    r += 1
    calc = [
        ("Assay floor L = 1000 / volume", f"=IF(B{volA}=\"\",\"\",1000/B{volA})",
         f"=IF(B{volB}=\"\",\"\",1000/B{volB})", "CFU/mL, one colony in the volume plated"),
        ("Seed needed for a q-log call, L*10^q",
         f"=IF(B{volA}=\"\",\"\",(1000/B{volA})*10^B{q})",
         f"=IF(B{volB}=\"\",\"\",(1000/B{volB})*10^B{q})",
         "the reachability boundary: below this the endpoint cannot be observed"),
        ("Identifiability boundary L/c1",
         f"=IF(OR(B{volA}=\"\",B{c1}=\"\"),\"\",(1000/B{volA})/B{c1})",
         f"=IF(OR(B{volB}=\"\",B{c1}=\"\"),\"\",(1000/B{volB})/B{c1})",
         "above this, a floored reading can only be called lowest class"),
        ("Headroom h at LOW inoculum",
         f"=IF(OR(B{lo}=\"\",B{volA}=\"\"),\"\",LOG10(B{lo}/(1000/B{volA})))",
         f"=IF(OR(B{lo}=\"\",B{volB}=\"\"),\"\",LOG10(B{lo}/(1000/B{volB})))",
         "log10(N0/L); the deepest endpoint this culture could ever show"),
        ("Headroom h at MID inoculum",
         f"=IF(OR(B{mid}=\"\",B{volA}=\"\"),\"\",LOG10(B{mid}/(1000/B{volA})))",
         f"=IF(OR(B{mid}=\"\",B{volB}=\"\"),\"\",LOG10(B{mid}/(1000/B{volB})))", ""),
        ("Headroom h at HIGH inoculum",
         f"=IF(OR(B{hi}=\"\",B{volA}=\"\"),\"\",LOG10(B{hi}/(1000/B{volA})))",
         f"=IF(OR(B{hi}=\"\",B{volB}=\"\"),\"\",LOG10(B{hi}/(1000/B{volB})))", ""),
    ]
    calc_first = r
    for what, fa, fb, meaning in calc:
        ws.cell(row=r, column=1, value=what).font = BODY
        for col, f in ((2, fa), (3, fb)):
            c = ws.cell(row=r, column=col, value=f)
            c.font, c.border = MONO, BOX
            c.fill = PatternFill("solid", fgColor=CALC_BG)
            c.number_format = "0.00"
        m = ws.cell(row=r, column=4, value=meaning)
        m.font, m.alignment = SMALL, WRAP
        r += 1

    hA_lo, hA_mid, hA_hi = calc_first + 3, calc_first + 4, calc_first + 5
    r += 1
    ws.cell(row=r, column=1, value="Is a q-log call legal? — the prediction, in one row").font = H2
    r += 1
    header(ws, r, ["Inoculum", "Plating A (100 uL)", "Plating B (10 uL)",
                   "If these disagree, the floor decided the label, not the drug"],
           [34, 18, 18, 62])
    r += 1
    for label, hrow in (("LOW", hA_lo), ("MID", hA_mid), ("HIGH", hA_hi)):
        ws.cell(row=r, column=1, value=label).font = BODY
        for col in (2, 3):
            c = ws.cell(row=r, column=col,
                        value=f"=IF({get_column_letter(col)}{hrow}=\"\",\"\","
                              f"IF({get_column_letter(col)}{hrow}>=B{q},"
                              f"\"legal\",\"IMPOSSIBLE\"))")
            c.font, c.border = MONO, BOX
            c.fill = PatternFill("solid", fgColor=CALC_BG)
            c.alignment = Alignment(horizontal="center")
        r += 1
    r += 1
    note(ws, r, "An 'IMPOSSIBLE' cell is the whole experiment: that culture cannot "
                "report the endpoint no matter how completely the drug works, and "
                "any tolerance call made from it is a reading of the inoculum and "
                "the pipette.")
    return ws


def prediction(wb):
    ws = wb.create_sheet("Prediction")
    ws.column_dimensions["A"].width = 108
    ws.cell(row=1, column=1, value="What is predicted, before any count exists").font = H1
    r = 3
    for para in [
        "Written from the Design sheet alone. Nothing here depends on a result, and "
        "nothing here should be edited once counting starts.",
        "",
        "1.  At the LOW inoculum, the deepest reduction that can be reported is the "
        "headroom h shown on the Design sheet, not the reduction the drug achieves. "
        "If h < q, no q-log call is possible in that arm, and every fully killed "
        "culture in it will read at the floor rather than below it.",
        "",
        "2.  For any sample whose count lands on the floor, the reported surviving "
        "fraction is L/N0 and nothing else. Two flasks killed to different true "
        "depths report the same fraction; two flasks killed identically but seeded "
        "differently report different fractions. The reported fraction therefore "
        "tracks the inoculum, not the killing.",
        "",
        "3.  The same sample plated at 100 uL and at 10 uL has two different floors "
        "and therefore two different reported fractions whenever the count is at or "
        "near the floor. Where the two platings straddle a class threshold, the "
        "SAME CULTURE receives two different tolerance labels. There is one culture, "
        "so the difference is not biology.",
        "",
        "4.  Above the identifiability boundary L/c1, a floored reading is compatible "
        "with the lowest class only. The label is then fixed by the seeding and the "
        "plated volume before the drug is added.",
        "",
        "WHAT WOULD REFUTE EACH. (1) fails if the low arm reports a q-log reduction. "
        "(2) fails if floored samples report fractions that track measured killing "
        "rather than N0. (3) fails if the two platings agree on every floored "
        "sample. (4) fails if a floored sample above L/c1 is compatible with more "
        "than one class under the stated thresholds.",
        "",
        "Signed, and dated before the first plate was counted:",
    ]:
        c = ws.cell(row=r, column=1, value=para)
        c.font, c.alignment = BODY, WRAP
        if para:
            ws.row_dimensions[r].height = 15 * (1 + len(para) // 95)
        r += 1
    r += 1
    for label in ("Name", "Date"):
        ws.cell(row=r, column=1, value=label).font = H2
        c = ws.cell(row=r, column=2)
        c.fill = PatternFill("solid", fgColor=ENTRY_BG)
        c.border = BOX
        ws.column_dimensions["B"].width = 26
        r += 2
    return ws


def counts(wb):
    ws = wb.create_sheet("Counts")
    ws.cell(row=1, column=1, value="Counts — one row per plate").font = H1
    note(ws, 2, "Yellow columns are entry. Blue columns compute. Leave 'colonies' "
                "empty only if the plate was not read; enter 0 for a genuinely "
                "blank plate, which is not the same thing.", span=12)
    hr = 4
    cols = ["Arm (inoculum)", "Drug conc (x MIC)", "Flask / replicate",
            "Day", "Plated volume (uL)", "Dilution factor (1 = neat)",
            "Colonies counted", "Plate note",
            "CFU/mL", "Floor L for this plating", "At or below floor?",
            "log10 CFU/mL"]
    header(ws, hr, cols, [16, 14, 14, 7, 14, 16, 14, 22, 14, 16, 14, 13])

    n_rows = 240
    for i in range(n_rows):
        r = hr + 1 + i
        for col in range(1, 9):
            c = ws.cell(row=r, column=col)
            c.fill = PatternFill("solid", fgColor=ENTRY_BG)
            c.border, c.font = BOX, BODY
        # CFU/mL = colonies * dilution * (1000 / volume plated)
        f = (f'=IF(OR(G{r}="",E{r}="",F{r}=""),"",G{r}*F{r}*1000/E{r})')
        ws.cell(row=r, column=9, value=f)
        ws.cell(row=r, column=10, value=f'=IF(E{r}="","",1000/E{r})')
        ws.cell(row=r, column=11,
                value=f'=IF(OR(I{r}="",J{r}=""),"",IF(I{r}<=J{r},"YES","no"))')
        ws.cell(row=r, column=12, value=f'=IF(OR(I{r}="",I{r}<=0),"",LOG10(I{r}))')
        for col in (9, 10, 11, 12):
            c = ws.cell(row=r, column=col)
            c.fill = PatternFill("solid", fgColor=CALC_BG)
            c.border, c.font = BOX, MONO
            c.number_format = "0.00" if col in (9, 10, 12) else "General"

    dv_arm = DataValidation(type="list", formula1='"LOW,MID,HIGH,untreated"',
                            allow_blank=True)
    dv_vol = DataValidation(type="list", formula1='"100,10"', allow_blank=True)
    ws.add_data_validation(dv_arm)
    ws.add_data_validation(dv_vol)
    dv_arm.add(f"A{hr+1}:A{hr+n_rows}")
    dv_vol.add(f"E{hr+1}:E{hr+n_rows}")
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    return ws


def day_zero(wb):
    ws = wb.create_sheet("Day 0 check")
    ws.cell(row=1, column=1, value="Did you get the inoculum you aimed for?").font = H1
    r = 2
    r = note(ws, r, "The boundaries use the REALISED starting density, not the "
                    "target. If the realised value differs, the experiment is not "
                    "spoiled — but the prediction must be re-read against what you "
                    "actually seeded, and that has to be visible.", span=6)
    r += 1
    header(ws, r, ["Arm", "Target N0 (CFU/mL)", "Realised N0 (CFU/mL, day 0)",
                   "Fold difference", "Realised headroom at 100 uL",
                   "Still supports a 4-log call?"], [12, 20, 26, 14, 24, 24])
    r += 1
    for arm in ("LOW", "MID", "HIGH"):
        ws.cell(row=r, column=1, value=arm).font = BODY
        for col in (2, 3):
            c = ws.cell(row=r, column=col)
            c.fill = PatternFill("solid", fgColor=ENTRY_BG)
            c.border, c.font = BOX, BODY
        ws.cell(row=r, column=4, value=f'=IF(OR(B{r}="",C{r}=""),"",C{r}/B{r})')
        ws.cell(row=r, column=5, value=f'=IF(C{r}="","",LOG10(C{r}/10))')
        ws.cell(row=r, column=6,
                value=f'=IF(E{r}="","",IF(E{r}>=4,"yes","NO — cannot report it"))')
        for col in (4, 5, 6):
            c = ws.cell(row=r, column=col)
            c.fill = PatternFill("solid", fgColor=CALC_BG)
            c.border, c.font = BOX, MONO
            c.number_format = "0.00" if col in (4, 5) else "General"
        r += 1
    r += 2
    note(ws, r, "Headroom here assumes the 100 uL plating, where L = 10 CFU/mL. "
                "The 10 uL plating has L = 100, so its headroom is one log smaller "
                "in every arm — which is the point of plating both ways.", span=6)
    return ws


def main() -> int:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    readme(wb)
    design(wb)
    prediction(wb)
    counts(wb)
    day_zero(wb)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")
    print(f"   sheets: {', '.join(wb.sheetnames)}")
    print(f"   {OUT.stat().st_size / 1024:.0f} KB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
